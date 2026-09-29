#!/usr/bin/env python3
"""Upload the current-period delivery workbook to Feishu by guarded 任务ID lookup.

The current workflow uses the local source-sheet `任务ID` as the workspace/delivery directory
name. The script re-reads the remote sheet, locates the row by `任务ID`, and
refuses to write unless the remote task title matches the local delivery
workbook's task title.

Default target is loaded from skills/qc-eval/config/g1.json.

Usage:
  python3 skills/qc-eval/scripts/push_g1_feishu.py <任务ID> --dry-run
  python3 skills/qc-eval/scripts/push_g1_feishu.py <任务ID> --write --confirm-user-instructed
"""

from __future__ import annotations

import argparse
import subprocess
import sys
import re
import time
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SKILL_DIR = HERE.parent
REPO_ROOT = SKILL_DIR.parent.parent
sys.path.insert(0, str(HERE))

from lib.g1 import (  # noqa: E402
    canonical_header_name,
    FEISHU_SHEET_ID,
    FEISHU_WIKI,
    PERIOD,
    is_empty_feishu_dirty_tag,
    is_feishu_dirty_tag_header,
    is_feishu_dirty_tag_related_header,
    to_feishu_dirty_tag,
)
from lib.label_options import LABEL_HEADER, label_options  # noqa: E402
from lib.concurrency import DirectoryLock  # noqa: E402
from lib.ai_accuracy import inspect_ai_accuracy_schema  # noqa: E402
from push_feishu import (  # noqa: E402
    READ_ONLY_HEADERS,
    build_push_targets,
    cell_to_text,
    clear_cell,
    get_sheet_meta,
    read_range,
    resolve_spreadsheet_token,
    write_cell,
    write_image,
)

DEFAULT_WIKI = FEISHU_WIKI
DEFAULT_SHEET_ID = FEISHU_SHEET_ID
DEFAULT_PROFILE = "quan"


def read_remote_headers(token: str, sheet_id: str, max_col: int, profile: str) -> dict[str, int]:
    last = get_column_letter(max_col)
    rows = read_range(token, sheet_id, f"A1:{last}3", profile)
    for row in rows:
        headers = {}
        for idx, value in enumerate(row, start=1):
            text = cell_to_text(value).strip()
            if text:
                headers[canonical_header_name(text)] = idx
        if "任务" not in headers and headers.get("任务ID", 0) > 1:
            headers["任务"] = headers["任务ID"] - 1
        if "任务" in headers and "任务ID" in headers:
            return headers
    raise RuntimeError("远端前 3 行没有找到「任务」和「任务ID」表头")


def safe_tag(name: str) -> str:
    return re.sub(r"[^0-9A-Za-z_.-]+", "_", name.strip())


def find_remote_row_by_task_id(
    token: str,
    sheet_id: str,
    row_count: int,
    headers: dict[str, int],
    task_id: str,
    profile: str,
) -> tuple[int | None, str]:
    first = min(headers["任务ID"], headers["任务"])
    last = max(headers["任务ID"], headers["任务"])
    values = read_range(
        token,
        sheet_id,
        f"{get_column_letter(first)}1:{get_column_letter(last)}{row_count}",
        profile,
    )
    id_idx = headers["任务ID"] - first
    task_idx = headers["任务"] - first
    for offset, row in enumerate(values, start=1):
        if id_idx < len(row) and cell_to_text(row[id_idx]).strip() == task_id:
            title = cell_to_text(row[task_idx]).strip() if task_idx < len(row) else ""
            return offset, title
    return None, ""


def read_remote_cell(token: str, sheet_id: str, cell: str, profile: str) -> str:
    """Return current remote cell text for fields that must not be overwritten."""
    rows = read_range(token, sheet_id, f"{cell}:{cell}", profile)
    if not rows or not rows[0]:
        return ""
    return cell_to_text(rows[0][0]).strip()


def load_delivery(task_id: str, local_xlsx: Path | None):
    xlsx = local_xlsx or (REPO_ROOT / "delivery" / task_id / f"{task_id}.xlsx")
    if not xlsx.exists():
        raise FileNotFoundError(f"本地交付 xlsx 不存在: {xlsx}")
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.worksheets[0]
    headers = {
        canonical_header_name(str(ws.cell(1, c).value)): c
        for c in range(1, ws.max_column + 1)
        if isinstance(ws.cell(1, c).value, str) and str(ws.cell(1, c).value).strip()
    }
    if "任务" not in headers:
        raise RuntimeError(f"{xlsx} 表头缺少「任务」列")
    return xlsx, ws, headers


def verify_screenshots(task_id: str, skip: bool) -> int:
    if skip:
        return 0
    guards = [
        "blank_guard.py",
        "evidence_guard.py",
        "screenshot_style_guard.py",
    ]
    for script in guards:
        rc = subprocess.call([sys.executable, str(HERE / script), task_id], cwd=str(REPO_ROOT))
        if rc != 0:
            print(f"==> push 中止：{script} 未通过")
            return rc
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("task_id", help=f"{PERIOD}源表「任务ID」，也是 tasks/ 和 delivery/ 子目录名")
    parser.add_argument("--wiki-url", default=DEFAULT_WIKI)
    parser.add_argument("--sheet-id", default=DEFAULT_SHEET_ID)
    parser.add_argument("--profile", default=DEFAULT_PROFILE)
    parser.add_argument("--local-xlsx", default=None)
    parser.add_argument("--dry-run", action="store_true", help="只打印计划写入，不实际写飞书")
    parser.add_argument("--write", action="store_true", help="实际写飞书；未指定时默认 dry-run")
    parser.add_argument("--verify-only", action="store_true", help="只读核对远端质检文字与本地交付是否一致")
    parser.add_argument(
        "--confirm-user-instructed",
        action="store_true",
        help="确认用户已明确要求上传/回写飞书；没有该确认时 --write 会被拒绝",
    )
    parser.add_argument("--only-changed-text", action="store_true", help="只写文字，跳过截图")
    parser.add_argument("--audit-reason-model", default="", help="审核理由汇总列仅替换该模型所在行，保留远端其他模型原文")
    parser.add_argument(
        "--only-headers",
        default="",
        help="只上传这些表头（逗号分隔），跳过截图和其他文字列",
    )
    parser.add_argument(
        "--only-image-model",
        default="",
        help="只上传该模型的验证截图，文字列和其他模型截图都不写",
    )
    parser.add_argument(
        "--list-remote-headers",
        action="store_true",
        help="只读打印远端表头及列号，不写入任何内容",
    )
    args = parser.parse_args()

    if args.sheet_id != DEFAULT_SHEET_ID:
        print(f"错误：当前流程只允许上传 sheet={DEFAULT_SHEET_ID}，拒绝 {args.sheet_id}", file=sys.stderr)
        return 2
    if "sheet=" in args.wiki_url and f"sheet={DEFAULT_SHEET_ID}" not in args.wiki_url:
        print(f"错误：当前流程只允许 URL sheet={DEFAULT_SHEET_ID}", file=sys.stderr)
        return 2
    if args.write and args.dry_run:
        print("错误：--write 和 --dry-run 不能同时使用", file=sys.stderr)
        return 2
    if args.write and not args.confirm_user_instructed:
        print(
            "错误：禁止主动上传飞书。只有用户明确要求上传/回写飞书后，才可使用 "
            "--write --confirm-user-instructed。",
            file=sys.stderr,
        )
        return 2

    dry = not args.write
    local_xlsx = Path(args.local_xlsx) if args.local_xlsx else None
    xlsx, ws, local_headers = load_delivery(args.task_id, local_xlsx)
    from check_ai_feedback import check_task as check_ai_feedback, workbook_labels
    issues = check_ai_feedback(REPO_ROOT / "tasks" / args.task_id, workbook_labels(ws))
    if issues:
        print("错误：AI评价证据未通过：" + "; ".join(issues), file=sys.stderr)
        return 2
    if "任务ID" in local_headers:
        local_task_id = str(ws.cell(2, local_headers["任务ID"]).value or "").strip()
        if local_task_id and local_task_id != args.task_id:
            print(f"错误：本地交付任务ID={local_task_id} 与参数 {args.task_id} 不一致", file=sys.stderr)
            return 2
    local_title = str(ws.cell(2, local_headers["任务"]).value or "").strip()
    if not local_title:
        print("错误：本地交付第 2 行「任务」为空", file=sys.stderr)
        return 2

    print(f"==> {PERIOD} Feishu push: task_id={args.task_id} dry-run={dry}")
    print(f"    local xlsx: {xlsx}")
    print(f"    local task: {local_title}")

    write_lock = None
    if args.write:
        lock_path = REPO_ROOT / ".qc-eval" / "locks" / f"feishu-{args.sheet_id}.lock"
        print(f"    waiting for Feishu write lock: {lock_path.relative_to(REPO_ROOT)}")
        write_lock = DirectoryLock(lock_path, timeout=600, stale_after=3600).acquire()
        print("    Feishu write lock acquired")

    token = resolve_spreadsheet_token(args.wiki_url, args.profile)
    meta = get_sheet_meta(token, args.sheet_id, args.profile)
    title = str(meta.get("title", "")).strip()
    print(f"    remote sheet title: {title}")
    row_count = int(meta["grid_properties"]["row_count"])
    remote_headers = read_remote_headers(
        token,
        args.sheet_id,
        int(meta["grid_properties"]["column_count"]),
        args.profile,
    )
    ai_schema = inspect_ai_accuracy_schema(local_headers, remote_headers)
    if ai_schema["errors"]:
        print(
            "错误：本地交付与飞书表的 AI 准确性字段结构不一致："
            + "；".join(ai_schema["errors"]),
            file=sys.stderr,
        )
        return 2
    if args.list_remote_headers:
        for header, col in sorted(remote_headers.items(), key=lambda item: item[1]):
            print(f"    {get_column_letter(col)}: {header!r}")
        return 0
    row, remote_title = find_remote_row_by_task_id(
        token, args.sheet_id, row_count, remote_headers, args.task_id, args.profile
    )
    if not row:
        print(f"错误：远端表未找到任务ID={args.task_id}", file=sys.stderr)
        return 2
    print(f"    remote task row {row}: {remote_title}")
    if remote_title != local_title:
        print("错误：远端任务ID对应的任务名与本地不一致，拒绝上传", file=sys.stderr)
        return 2

    only_image_model = str(args.only_image_model or "").strip()
    if args.only_headers and only_image_model:
        print("错误：--only-headers 和 --only-image-model 不能同时使用", file=sys.stderr)
        return 2
    if verify_screenshots(
        args.task_id,
        args.only_changed_text or (bool(args.only_headers) and not only_image_model),
    ) != 0:
        return 1

    text_headers, image_map = build_push_targets(local_headers)
    if LABEL_HEADER in local_headers and LABEL_HEADER in remote_headers and LABEL_HEADER not in text_headers:
        text_headers.insert(0, LABEL_HEADER)
    only_headers = [h.strip() for h in str(args.only_headers or "").split(",") if h.strip()]
    if only_headers:
        missing = [h for h in only_headers if h not in local_headers or h not in remote_headers]
        if missing:
            print(f"错误：--only-headers 找不到列: {', '.join(missing)}", file=sys.stderr)
            return 2
        text_headers = [h for h in only_headers if h in text_headers]
        image_map = {}
    if only_image_model:
        text_headers = []
        matched = {
            header: model
            for header, model in image_map.items()
            if model == only_image_model
            or only_image_model in model
            or model.endswith(only_image_model)
        }
        if not matched:
            print(
                f"错误：--only-image-model 找不到模型 {only_image_model!r}。"
                f"可选: {', '.join(sorted(set(image_map.values())))}",
                file=sys.stderr,
            )
            return 2
        image_map = matched
        print(f"    只上传截图: {', '.join(image_map.values())}")
    text_count = 0
    if args.verify_only:
        cols = [remote_headers[h] for h in text_headers if h in remote_headers]
        values = read_range(token, args.sheet_id, f"A{row}:{get_column_letter(max(cols))}{row}", args.profile)
        remote_values = values[0] if values else []
        mismatches = 0
        for header in text_headers:
            if header in READ_ONLY_HEADERS or header == LABEL_HEADER:
                continue
            lcol, rcol = local_headers.get(header), remote_headers.get(header)
            if not lcol or not rcol:
                continue
            expected = str(ws.cell(2, lcol).value or "").strip()
            actual = cell_to_text(remote_values[rcol-1] if rcol <= len(remote_values) else None).strip()
            if is_feishu_dirty_tag_header(header):
                expected = "" if is_empty_feishu_dirty_tag(expected) else to_feishu_dirty_tag(expected)
            if actual != expected:
                mismatches += 1
                print(f"     [mismatch] {get_column_letter(rcol)}{row} {header}: expected={expected!r} actual={actual!r}")
        print(f"==> 只读核对完成：{mismatches} 个不一致单元格")
        return 1 if mismatches else 0
    for header in text_headers:
        if header in READ_ONLY_HEADERS:
            continue
        lcol = local_headers.get(header)
        rcol = remote_headers.get(header)
        if not lcol or not rcol:
            continue
        value = ws.cell(2, lcol).value
        text = "" if value is None else str(value)
        cell_a1 = f"{get_column_letter(rcol)}{row}"
        if args.audit_reason_model and header == "一次质检平台审核理由（按模型）":
            prefix = args.audit_reason_model + "（"
            replacement = [line for line in text.splitlines() if line.startswith(prefix)]
            remote_text = read_remote_cell(token, args.sheet_id, cell_a1, args.profile)
            lines = remote_text.splitlines(keepends=True)
            matches = [i for i, line in enumerate(lines) if line.startswith(prefix)]
            if len(replacement) != 1 or len(matches) != 1:
                raise ValueError("审核理由无法唯一匹配模型，停止写入汇总列")
            index = matches[0]
            ending = lines[index][len(lines[index].rstrip('\r\n')):]
            lines[index] = replacement[0] + ending
            text = ''.join(lines)
        if header == LABEL_HEADER:
            remote_text = read_remote_cell(token, args.sheet_id, cell_a1, args.profile)
            if remote_text and remote_text in set(label_options()):
                print(f"     [skip] {cell_a1} label already set: {remote_text!r}")
                continue
            if remote_text:
                print(f"     [fix] {cell_a1} replacing invalid label: {remote_text!r}")
            if not text.strip():
                continue
        if is_feishu_dirty_tag_header(header):
            # Dropdown cells: writing "" is not a valid option and Feishu shows
            # 「数据无效」. Clear content instead, same as a manual cell clear.
            if is_empty_feishu_dirty_tag(text):
                clear_cell(token, args.sheet_id, cell_a1, args.profile, dry)
            else:
                write_cell(
                    token,
                    args.sheet_id,
                    cell_a1,
                    to_feishu_dirty_tag(text),
                    args.profile,
                    dry,
                )
            text_count += 1
            continue
        if not text.strip():
            # Reinspection can withdraw a redline or restore verifiability.
            # Empty QC-owned reasons must clear the previous remote verdict.
            if is_feishu_dirty_tag_related_header(header) or header in {
                "不可验证原因", "红线问题", "题目红线问题", "题目红线备注",
            }:
                clear_cell(token, args.sheet_id, cell_a1, args.profile, dry)
                if not dry:
                    # Feishu reads may briefly retain the value after a successful clear.
                    for delay in (0, 0.5, 1, 2, 4):
                        if delay:
                            time.sleep(delay)
                        if not read_remote_cell(token, args.sheet_id, cell_a1, args.profile):
                            break
                    else:
                        raise ValueError(f"{cell_a1} 清除后回读仍有内容（已重试）")
                text_count += 1
            continue
        write_cell(token, args.sheet_id, cell_a1, text, args.profile, dry)
        if only_headers and not dry:
            if read_remote_cell(token, args.sheet_id, cell_a1, args.profile) != text:
                raise ValueError(f"{cell_a1} 写入后回读不一致")
            print(f"     [verified] {cell_a1}")
        text_count += 1

    image_count = 0
    if only_image_model or (not args.only_changed_text and not only_headers):
        img_dir = REPO_ROOT / "delivery" / args.task_id / "screenshots"
        for header, model in image_map.items():
            rcol = remote_headers.get(header)
            if not rcol:
                continue
            img = img_dir / f"{model.replace(' ', '_')}_verify.png"
            if not img.exists():
                img = img_dir / f"{safe_tag(model)}_verify.png"
            if not img.exists():
                print(f"   [warn] 本地缺图: {img.relative_to(REPO_ROOT)}")
                continue
            write_image(token, args.sheet_id, f"{get_column_letter(rcol)}{row}", img, args.profile, dry)
            image_count += 1

    print(f"==> 完成：文字 {text_count} 单元格，图片 {image_count} 单元格 ({'dry-run' if dry else '已写入'})")
    if write_lock:
        write_lock.release()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
