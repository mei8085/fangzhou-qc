#!/usr/bin/env python3
"""Mirror the 38期 Feishu QC sheet into a local xlsx snapshot.

This keeps the rest of the qc-eval workflow table-driven while the official
working sheet lives in Feishu.

Default input:
  https://bytedance.larkoffice.com/wiki/OY1Dw0dcailCxTk3UGOcykVznAf?sheet=uBlwx4

Default output:
  D:/fz/fangzhou/众测质检38.xlsx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from push_feishu import (  # noqa: E402
    cell_to_text,
    get_sheet_meta,
    read_range,
    resolve_spreadsheet_token,
)

DEFAULT_WIKI = "https://bytedance.larkoffice.com/wiki/OY1Dw0dcailCxTk3UGOcykVznAf?sheet=uBlwx4"
DEFAULT_SHEET_ID = "uBlwx4"
DEFAULT_PROFILE = "quan"
DEFAULT_OUT = Path("D:/fz/fangzhou/众测质检38.xlsx")
PERIOD_CHOICES = {
    "38期": (DEFAULT_WIKI, DEFAULT_SHEET_ID),
}


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--period", choices=sorted(PERIOD_CHOICES),
                   help="使用内置期数配置；本工程只允许 38期")
    p.add_argument("--wiki-url", default=DEFAULT_WIKI)
    p.add_argument("--sheet-id", default=DEFAULT_SHEET_ID)
    p.add_argument("--profile", default=DEFAULT_PROFILE)
    p.add_argument("--out", default=str(DEFAULT_OUT))
    args = p.parse_args()
    if args.period:
        args.wiki_url, args.sheet_id = PERIOD_CHOICES[args.period]
    if args.sheet_id != DEFAULT_SHEET_ID:
        print(f"ERROR: 本工程只允许同步 38期 sheet={DEFAULT_SHEET_ID}，拒绝 sheet={args.sheet_id}", file=sys.stderr)
        return 2
    if "sheet=" in args.wiki_url and f"sheet={DEFAULT_SHEET_ID}" not in args.wiki_url:
        print(f"ERROR: 本工程只允许同步 38期 URL：{DEFAULT_WIKI}", file=sys.stderr)
        return 2

    print(f"==> wiki url: {args.wiki_url}")
    token = resolve_spreadsheet_token(args.wiki_url, args.profile)
    meta = get_sheet_meta(token, args.sheet_id, args.profile)
    title = meta["title"]
    rows = meta["grid_properties"]["row_count"]
    cols = meta["grid_properties"]["column_count"]
    print(f"    sheet = {title} ({args.sheet_id}) rows={rows} cols={cols}")

    last = get_column_letter(cols)
    values = read_range(token, args.sheet_id, f"A1:{last}{rows}", args.profile)
    if not values or not any(cell_to_text(v).strip() in {"任务", "任务ID"} for v in values[0]):
        print("ERROR: 远端第 1 行不像 38 期质检表头，未写本地 xlsx", file=sys.stderr)
        return 1

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    for r_idx, row in enumerate(values, start=1):
        for c_idx, v in enumerate(row, start=1):
            ws.cell(r_idx, c_idx).value = cell_to_text(v)

    # Light usability formatting. Delivery xlsx generation remains header-driven.
    ws.freeze_panes = "A2"
    for c in range(1, cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"==> wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
