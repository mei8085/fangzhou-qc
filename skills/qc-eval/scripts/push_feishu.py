#!/usr/bin/env python3
"""Upload a task's local delivery into the 38期 Feishu spreadsheet — text
cells and verification screenshots only, scoped to the row whose 任务ID matches.

The script never runs by itself: it's invoked explicitly with one or more task
IDs by the user, after they've eyeballed `delivery/<TASK_ID>/`. This repository
is locked to the 38期coding质检 sheet inside the wiki document at:

    https://bytedance.larkoffice.com/wiki/OY1Dw0dcailCxTk3UGOcykVznAf?sheet=uBlwx4

It uses `lark-cli` with the `quan` profile. Network calls run with
`LARK_CLI_NO_PROXY=1` to bypass the local HTTP proxy (otherwise wiki/spaces
returns 404).

Safety guarantees:
  * Each task row is located **on the remote side** by matching column B
    (`任务ID`). Local row numbers from the source xlsx are NOT trusted: the
    remote sheet has different ordering.
  * Only reviewer-owned text fields and verification screenshot cells are
    touched. Every other column on that row, every other row, every other
    sheet — untouched. The script refuses any target sheet except uBlwx4.
  * `--dry-run` prints the exact cell + payload for each planned write without
    calling the write API.

Usage:
    push_feishu.py <TASK_ID> [<TASK_ID> ...] [--dry-run]
                   [--wiki-url URL] [--sheet-id SHEET_ID] [--profile quan]
                   [--only-changed-text] [--only-image-model MODEL]
                   [--confirm-user-instructed]

`--only-changed-text` 场景：题目已经推过一次，只是文字返修（比如换了评价口径），
截图不用换。开启后：批量读远端行的文字单元格 → 逐格 diff → 只写差异的文字格 →
完全跳过截图上传；blank_guard 也一并跳过（本来就没用到截图）。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import shutil
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SKILL_DIR = HERE.parent
REPO_ROOT = SKILL_DIR.parent.parent

DEFAULT_WIKI = "https://bytedance.larkoffice.com/wiki/OY1Dw0dcailCxTk3UGOcykVznAf?sheet=uBlwx4"
DEFAULT_SHEET_ID = "uBlwx4"  # 38期 sheet
DEFAULT_PROFILE = "quan"

# This repository is a 38期 workspace. Keep a single Feishu target to avoid
# accidentally writing another sheet.
PERIOD_TARGETS = {
    "38期": (DEFAULT_WIKI, DEFAULT_SHEET_ID),
}
ALLOWED_REMOTE_TITLE_PREFIX = "38期"


def sheet_id_from_wiki_url(url: str) -> str | None:
    """Use ?sheet=... from a pasted Feishu wiki URL when present.

    This prevents accidental uploads to a non-38期 sheet when a user
    provides a full link for another period.
    """
    try:
        values = parse_qs(urlparse(url).query).get("sheet") or []
    except Exception:
        return None
    return values[0] if values else None


def infer_period_from_delivery(task_ids: list[str]) -> str | None:
    """Infer period from delivery/<TASK_ID>/<TASK_ID>.xlsx worksheet titles.

    Returns a period like "24期" only when all existing delivery workbooks agree.
    Missing workbooks are ignored; callers still validate files later.
    """
    periods: set[str] = set()
    for task_id in task_ids:
        xlsx = REPO_ROOT / "delivery" / task_id / f"{task_id}.xlsx"
        if not xlsx.exists():
            continue
        try:
            wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
            title = wb.worksheets[0].title
        except Exception:
            continue
        clean = title.strip()
        if clean == "38期coding质检":
            periods.add("38期")
            continue
        if clean.startswith("38期"):
            periods.add("38期")
            continue
        m = re.match(r"(\d+期)", clean)
        if m:
            periods.add(m.group(1))
    if not periods:
        return None
    if len(periods) > 1:
        raise RuntimeError(f"多个任务交付物期数不一致: {sorted(periods)}")
    return next(iter(periods))


def choose_target(task_ids: list[str], wiki_url: str | None, sheet_id: str | None) -> tuple[str, str, str | None]:
    """Resolve target wiki/sheet.

    The delivery workbook title is only reported for diagnostics. Uploads from
    this repo are always forced to the 38期 sheet, and explicit overrides to
    any other sheet are rejected.
    """
    explicit_sheet = sheet_id or (sheet_id_from_wiki_url(wiki_url) if wiki_url else None)
    if wiki_url or explicit_sheet:
        target_sheet = explicit_sheet or DEFAULT_SHEET_ID
        if target_sheet != DEFAULT_SHEET_ID:
            raise RuntimeError(
                f"本工程只允许上传 38期 sheet={DEFAULT_SHEET_ID}，拒绝目标 sheet={target_sheet}"
            )
        return DEFAULT_WIKI, DEFAULT_SHEET_ID, infer_period_from_delivery(task_ids)

    period = infer_period_from_delivery(task_ids)
    return DEFAULT_WIKI, DEFAULT_SHEET_ID, period


def pushed_to_allowed_sheet(sentinel: Path) -> bool:
    if not sentinel.exists():
        return False
    try:
        text = sentinel.read_text(encoding="utf-8")
    except Exception:
        return False
    return f"sheet_id={DEFAULT_SHEET_ID}" in text

# 期数无关的固定文字列（每期都有）。
_TEXT_HEADER_ALWAYS = [
    "是否是高质量题目",
    "任务质量（题目质量）",
    "题目质量说明",
    "题目是否可验证",
    "不可验证原因",
    "分类是否正确",
]
_REDLINE_HEADERS = ["红线问题", "题目红线问题", "题目红线备注"]
_TEXT_HEADER_TAIL = ["质检人"]
_NORMAL_TASK_HEADERS = ["是否正常题目"]
_SUPPLIER_HEADERS = ["所属商", "所述商"]
# 期数相关的可选文字列（有就写、没就跳）。
_TEXT_HEADER_OPTIONAL = [
    "一次质检平台审核理由（按模型）",
    "一次质检是否完成",
    "0614C对比GLM 5.2 模型差异（轨迹/产物）pattern简述",
    "0614C完成质量实际打分",
    "GLM5.2完成质量打分",
    "机器质检结果是否正确",
    "错误原因",
]
_MODEL_TEXT_SUFFIXES = ["反馈质量（作业质量）", "质量备注", "脏数据标签", "脏数据标签理由"]
_MODEL_IMAGE_SUFFIX = "验证截图"
READ_ONLY_HEADERS = {"是否结算"}


def build_push_targets(local_headers: dict[str, int]) -> tuple[list[str], dict[str, str]]:
    """从 delivery xlsx 表头动态推断 (文字列 headers, 截图列 header→model 名)。

    模型名从表头里所有 `<model>反馈质量（作业质量）` 项抽出，保证 28/29 期
    都能自适应。
    """
    text_headers = list(_TEXT_HEADER_ALWAYS)
    text_headers.extend(h for h in _REDLINE_HEADERS if h in local_headers)
    text_headers.extend(h for h in _TEXT_HEADER_TAIL if h in local_headers)
    text_headers.extend(h for h in _NORMAL_TASK_HEADERS if h in local_headers)
    text_headers.extend(h for h in _SUPPLIER_HEADERS if h in local_headers)
    text_headers.extend(h for h in _TEXT_HEADER_OPTIONAL if h in local_headers)

    model_names: list[str] = []
    for h in local_headers:
        if h.endswith(_MODEL_TEXT_SUFFIXES[0]):
            model_names.append(h[:-len(_MODEL_TEXT_SUFFIXES[0])])

    for model in model_names:
        for suf in _MODEL_TEXT_SUFFIXES:
            text_headers.append(f"{model}{suf}")
    text_headers = [h for h in text_headers if h not in READ_ONLY_HEADERS]

    image_map = {f"{model}{_MODEL_IMAGE_SUFFIX}": model for model in model_names}
    return text_headers, image_map


# ---------- lark-cli wrapper ----------

_LARK_EXE: list[str] | None = None


def lark_exe() -> list[str]:
    """Return the subprocess argv prefix that runs lark-cli on this OS.

    On Windows the user-facing `lark-cli` in PATH is a bash wrapper that
    Python's subprocess can't launch directly, so we resolve the underlying
    `.exe` once and cache it. On macOS/Linux we prefer `shutil.which("lark-cli")`.
    """
    global _LARK_EXE
    if _LARK_EXE is not None:
        return _LARK_EXE
    candidates: list[list[str]] = []
    which = shutil.which("lark-cli")
    if which:
        candidates.append([which])
    if sys.platform == "win32":
        cmd_path = shutil.which("lark-cli.cmd")
        if cmd_path:
            cmd_dir = Path(cmd_path).resolve().parent
            exe_path = cmd_dir / "node_modules" / "@larksuite" / "cli" / "bin" / "lark-cli.exe"
            if exe_path.exists():
                candidates.append([str(exe_path)])
        candidates.extend([["lark-cli.exe"], ["lark-cli.cmd"], ["cmd", "/c", "lark-cli"]])
        roots = [
            Path("C:/nvm4w/nodejs/node_modules/@larksuite"),
            Path(os.path.expandvars("%APPDATA%/npm/node_modules/@larksuite")),
        ]
        for root in roots:
            if not root.exists():
                continue
            for exe in root.glob("*/bin/lark-cli.exe"):
                candidates.append([str(exe)])
    else:
        candidates.append(["lark-cli"])
        npm_cmd = shutil.which("npm")
        if npm_cmd:
            try:
                npm_root = subprocess.run(
                    [npm_cmd, "root", "-g"],
                    capture_output=True, text=True, check=False,
                ).stdout.strip()
            except OSError:
                npm_root = ""
            if npm_root:
                cli_js = Path(npm_root) / "@larksuite" / "cli" / "bin" / "lark-cli.js"
                if cli_js.exists():
                    node = shutil.which("node") or "node"
                    candidates.append([node, str(cli_js)])
    seen: set[tuple[str, ...]] = set()
    unique: list[list[str]] = []
    for candidate in candidates:
        key = tuple(candidate)
        if key in seen:
            continue
        seen.add(key)
        unique.append(candidate)
    for candidate in unique:
        try:
            r = subprocess.run([*candidate, "--version"], capture_output=True, text=True)
            if r.returncode == 0:
                _LARK_EXE = candidate
                return _LARK_EXE
        except (FileNotFoundError, OSError):
            pass
    raise RuntimeError("lark-cli 可执行文件没找到；确认已安装且在 PATH 上")


def lark(args: list[str], *, profile: str, dry: bool = False, capture: bool = True, cwd: str | None = None) -> dict | str | None:
    cmd = [*lark_exe(), "--profile", profile, *args]
    env = os.environ.copy()
    env["LARK_CLI_NO_PROXY"] = "1"
    if dry:
        # show a redacted command (image flag stays, write-image needs to be visible)
        shown = " ".join(cmd[:7]) + ("..." if len(cmd) > 7 else "")
        print(f"     [dry] $ {shown}")
        return None
    r = subprocess.run(cmd, capture_output=True, text=True, env=env, cwd=cwd)
    if r.returncode != 0:
        raise RuntimeError(f"lark-cli error ({r.returncode}): {r.stderr.strip() or r.stdout.strip()}")
    out = r.stdout.strip()
    if not capture:
        return None
    if not out:
        return {}
    try:
        return json.loads(out)
    except json.JSONDecodeError:
        return out


def resolve_spreadsheet_token(wiki_url: str, profile: str) -> str:
    m = re.search(r"/wiki/([A-Za-z0-9]+)", wiki_url)
    if not m:
        raise ValueError(f"无法从 URL 提取 wiki node token: {wiki_url}")
    node_token = m.group(1)
    d = lark(
        ["wiki", "spaces", "get_node", "--params", json.dumps({"token": node_token})],
        profile=profile,
    )
    node = d["data"]["node"]
    if node["obj_type"] != "sheet":
        raise RuntimeError(f"wiki 节点不是 sheet (是 {node['obj_type']})")
    return node["obj_token"]


def read_range(token: str, sheet_id: str, a1: str, profile: str, value_render: str = "ToString") -> list[list]:
    d = lark(
        ["sheets", "+read",
         "--spreadsheet-token", token, "--sheet-id", sheet_id,
         "--range", a1, "--value-render-option", value_render],
        profile=profile,
    )
    return d.get("data", {}).get("valueRange", {}).get("values") or []


def read_row_cells(token: str, sheet_id: str, row: int, cols: list[int], profile: str) -> dict[int, str]:
    """Batch-read a set of cells in the same row and return {col_index: text_value}.

    Reads A:<last_col> for that single row so we get everything in one round-trip,
    then picks out the requested column indices. Column indices are 1-based.
    """
    if not cols:
        return {}
    last = get_column_letter(max(cols))
    first = "A"
    rows = read_range(token, sheet_id, f"{first}{row}:{last}{row}", profile)
    row_vals = rows[0] if rows else []
    out: dict[int, str] = {}
    for c in cols:
        v = row_vals[c - 1] if c - 1 < len(row_vals) else None
        out[c] = cell_to_text(v).strip()
    return out


def write_cell(token: str, sheet_id: str, cell: str, value, profile: str, dry: bool) -> None:
    # +cells-set takes a 2D JSON array of cell objects; a single cell is
    # [[{"value": value}]].
    payload = json.dumps([[{"value": value}]], ensure_ascii=False)
    args = ["sheets", "+cells-set",
            "--spreadsheet-token", token, "--sheet-id", sheet_id,
            "--range", cell, "--cells", payload,
            "--allow-overwrite"]
    if dry:
        preview = str(value)
        preview = preview[:60] + ("…" if len(preview) > 60 else "")
        print(f"     [dry] write {cell} = {preview!r}")
        return
    lark(args, profile=profile)


def write_image(token: str, sheet_id: str, cell: str, image: Path, profile: str, dry: bool) -> None:
    # lark-cli 拒收绝对路径，强制走 image 的父目录 + basename
    args = ["sheets", "+cells-set-image",
            "--spreadsheet-token", token, "--sheet-id", sheet_id,
            "--range", cell, "--image", f"./{image.name}"]
    if dry:
        print(f"     [dry] write-image {cell} ← {image.name} ({image.stat().st_size//1024} KiB)")
        return
    lark(args, profile=profile, cwd=str(image.parent))


# ---------- helpers ----------

def cell_to_text(v) -> str:
    """Normalize a sheet cell (string | rich-text array | None) → plain string."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, list):
        return "".join(seg.get("text", "") if isinstance(seg, dict) else str(seg) for seg in v)
    return str(v)


def read_header_columns(token: str, sheet_id: str, max_col: int, profile: str) -> dict[str, int]:
    """Detect the header row (1..3) and return {header_text: col_index_1_based}.

    28/29 期导出表常见 row 2 表头，33/34 期 Feishu 正式表是 row 1 表头。
    """
    last_letter = get_column_letter(max_col)
    rows = read_range(token, sheet_id, f"A1:{last_letter}3", profile)
    if not rows:
        raise RuntimeError("无法读取远端表头行")
    for row in rows:
        out = {}
        for idx, v in enumerate(row, start=1):
            s = cell_to_text(v).strip()
            if s:
                out[s] = idx
        if "任务ID" in out:
            repair_model_rating_headers(out)
            return out
    raise RuntimeError("远端前 3 行没有找到 任务ID 表头")


def repair_model_rating_headers(headers: dict[str, int]) -> None:
    """Repair known malformed model rating headers in-place.

    The 24 期 Feishu sheet currently has column M header accidentally set to
    "1", while N/O are still "glm5.1验证截图" and "glm5.1质量备注". Downstream
    matching is header-driven, so infer the missing `<model>反馈质量（作业质量）`
    column as the cell immediately before `<model>验证截图`.
    """
    for h, col in list(headers.items()):
        if not h.endswith(_MODEL_IMAGE_SUFFIX):
            continue
        model = h[:-len(_MODEL_IMAGE_SUFFIX)]
        rating_header = f"{model}{_MODEL_TEXT_SUFFIXES[0]}"
        note_header = f"{model}{_MODEL_TEXT_SUFFIXES[1]}"
        if rating_header in headers or note_header not in headers:
            continue
        if headers[note_header] == col + 1:
            headers[rating_header] = col - 1
            print(f"    [repair] infer header {rating_header!r} at col {col - 1}")


def find_remote_row_by_task_id(token: str, sheet_id: str, id_col: int, task_id: str,
                               row_count: int, profile: str) -> int | None:
    """Read the 任务ID column and return the 1-based row where it equals task_id."""
    col = get_column_letter(id_col)
    rows = read_range(token, sheet_id, f"{col}1:{col}{row_count}", profile)
    for i, r in enumerate(rows, start=1):
        if not r:
            continue
        if cell_to_text(r[0]).strip() == task_id:
            return i
    return None


def get_sheet_meta(token: str, sheet_id: str, profile: str) -> dict:
    """Return the requested sheet's row_count/column_count via +info."""
    d = lark(["sheets", "+info", "--spreadsheet-token", token], profile=profile)
    for s in d["data"]["sheets"]["sheets"]:
        if s["sheet_id"] == sheet_id:
            return s
    raise RuntimeError(f"sheet_id={sheet_id} 在该 spreadsheet 里不存在")


# ---------- per-task push ----------

def push_one(task_id: str, token: str, sheet_id: str, remote_headers: dict[str, int],
             row_count: int, profile: str, dry: bool,
             only_changed_text: bool = False,
             only_image_model: str | None = None) -> tuple[int, int]:
    print(f"\n==> {task_id}")

    # Locate local delivery xlsx
    local_xlsx = REPO_ROOT / "delivery" / task_id / f"{task_id}.xlsx"
    if not local_xlsx.exists():
        print(f"   [skip] 本地 {local_xlsx.relative_to(REPO_ROOT)} 不存在")
        return 0, 0
    wb = openpyxl.load_workbook(local_xlsx, data_only=True)
    ws = wb.worksheets[0]
    local_headers = {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1)
                     if isinstance(ws.cell(1, c).value, str) and str(ws.cell(1, c).value).strip()}

    # 从 delivery xlsx 表头动态推断本次要 push 的列（跨期通用）
    text_headers, image_map = build_push_targets(local_headers)

    # Locate remote row by task ID
    id_col = remote_headers.get("任务ID")
    if not id_col:
        print(f"   [err] 远端表头没找到 任务ID")
        return 0, 0
    row = find_remote_row_by_task_id(token, sheet_id, id_col, task_id, row_count, profile)
    if not row:
        print(f"   [skip] 远端 sheet 里没找到 任务ID = {task_id}")
        return 0, 0

    # Sanity: title match (best-effort, just warning)
    title_col_r = remote_headers.get("任务")
    title_col_l = local_headers.get("任务")
    if title_col_r and title_col_l:
        remote_title = cell_to_text(
            (read_range(token, sheet_id, f"{get_column_letter(title_col_r)}{row}", profile)[0:1] or [[None]])[0][0]
        ).strip()
        local_title = str(ws.cell(2, title_col_l).value or "").strip()
        if remote_title and remote_title != local_title:
            print(f"   [warn] 远端 row {row} 题目=「{remote_title}」 vs 本地「{local_title}」，仍按 任务ID 继续")
    print(f"   远端 row = {row}")

    # 计划写入的文字列 (header, local_col, remote_col, local_value)
    text_plan: list[tuple[str, int, int, str]] = []
    for header in text_headers:
        if header in READ_ONLY_HEADERS:
            print(f"   [skip] 只读公式列不上传: {header}")
            continue
        lcol = local_headers.get(header)
        rcol = remote_headers.get(header)
        if not lcol:
            print(f"   [warn] 本地缺列: {header}")
            continue
        if not rcol:
            print(f"   [warn] 远端缺列: {header}")
            continue
        value = ws.cell(2, lcol).value
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        text_plan.append((header, lcol, rcol, str(value)))

    # --only-changed-text 模式：批量读远端行的目标单元格，只保留差异
    if only_changed_text and text_plan:
        remote_vals = read_row_cells(
            token, sheet_id, row, [rcol for _, _, rcol, _ in text_plan], profile
        )
        filtered: list[tuple[str, int, int, str]] = []
        for header, lcol, rcol, local_v in text_plan:
            remote_v = remote_vals.get(rcol, "")
            if local_v.strip() == remote_v.strip():
                continue
            filtered.append((header, lcol, rcol, local_v))
        skipped = len(text_plan) - len(filtered)
        if skipped:
            print(f"   [diff] 跳过 {skipped} 格未变化的文字列")
        text_plan = filtered

    # Push text cells
    txt_count = 0
    if only_image_model:
        print(f"   [image-only] 文字列整体跳过（只更新 {only_image_model} 截图）")
    else:
        for header, _lcol, rcol, value in text_plan:
            cell_a1 = f"{get_column_letter(rcol)}{row}"
            write_cell(token, sheet_id, cell_a1, value, profile, dry)
            txt_count += 1

    # Push images (--only-changed-text 模式完全跳过截图)
    img_count = 0
    if only_changed_text:
        print(f"   [diff] 截图列整体跳过（--only-changed-text）")
    else:
        img_dir = REPO_ROOT / "delivery" / task_id / "screenshots"
        for header, model in image_map.items():
            if only_image_model and model != only_image_model:
                continue
            rcol = remote_headers.get(header)
            if not rcol:
                print(f"   [warn] 远端缺列: {header}")
                continue
            img = img_dir / f"{model.replace(' ', '_')}_verify.png"
            if not img.exists():
                print(f"   [warn] 本地缺图: {img.relative_to(REPO_ROOT)}")
                continue
            cell_a1 = f"{get_column_letter(rcol)}{row}"
            write_image(token, sheet_id, cell_a1, img, profile, dry)
            img_count += 1

    print(f"   ✓ {task_id} 文字 {txt_count} 单元格，图片 {img_count} 单元格 ({'dry-run' if dry else '已写入'})")
    return txt_count, img_count


# ---------- main ----------

def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("task_ids", nargs="+", help="一个或多个任务 ID")
    p.add_argument("--wiki-url", default=None)
    p.add_argument("--sheet-id", default=None)
    p.add_argument("--profile", default=DEFAULT_PROFILE)
    p.add_argument("--dry-run", action="store_true",
                   help="不实际写飞书，只打印将要执行的单元格更新")
    p.add_argument(
        "--confirm-user-instructed",
        action="store_true",
        help="确认用户已明确要求上传/回写飞书；没有该确认时非 dry-run 写入会被拒绝",
    )
    p.add_argument("--only-changed-text", action="store_true",
                   help="返修专用：批量读远端行 → 逐格 diff → 只写差异的文字格；截图完全跳过")
    p.add_argument("--only-image-model",
                   help="返修专用：只更新指定模型名的验证截图，文字列不写，其它截图不写")
    args = p.parse_args()
    if not args.dry_run and not args.confirm_user_instructed:
        print(
            "错误：禁止主动上传飞书。只有用户明确要求上传/回写/写入飞书后，才可使用 "
            "--confirm-user-instructed 执行非 dry-run 写入。",
            file=sys.stderr,
        )
        return 2
    try:
        wiki_url, sheet_id, inferred_period = choose_target(args.task_ids, args.wiki_url, args.sheet_id)
    except RuntimeError as e:
        print(f"错误：{e}", file=sys.stderr)
        return 2

    print(f"==> wiki url: {wiki_url}")
    print(f"    sheet id: {sheet_id}   profile: {args.profile}   dry-run: {args.dry_run}")
    if inferred_period:
        print(f"    inferred period: {inferred_period}")
        if inferred_period != "38期":
            print(f"    [warn] 交付物标题显示 {inferred_period}，但本工程强制上传 38期 uBlwx4")

    print("==> 解析 spreadsheet token")
    token = resolve_spreadsheet_token(wiki_url, args.profile)
    print(f"    spreadsheet_token = {token}")

    print("==> 读取 sheet meta")
    meta = get_sheet_meta(token, sheet_id, args.profile)
    print(f"    title = {meta['title']}  rows = {meta['grid_properties']['row_count']}"
          f"  cols = {meta['grid_properties']['column_count']}")
    remote_title = str(meta.get("title", "")).strip()
    if sheet_id != DEFAULT_SHEET_ID or not remote_title.startswith(ALLOWED_REMOTE_TITLE_PREFIX):
        print(
            f"错误：远端目标不是 38期，sheet_id={sheet_id} title={meta.get('title')!r}，已中止",
            file=sys.stderr,
        )
        return 2

    print("==> 读取远端表头")
    remote_headers = read_header_columns(token, sheet_id, meta["grid_properties"]["column_count"], args.profile)
    print(f"    远端列数 = {len(remote_headers)}")

    # 截图卡点：扫 delivery/<ID>/screenshots/，任一帧接近纯色、缺少实际运行证据、
    # 或验证图带旧式合成背景/边框就拒推。dry-run 也做检查，避免临上传才发现问题。
    # --only-changed-text 模式不推截图，也跳过截图 guards。
    if not args.only_changed_text:
        guards = [
            (
                "blank_guard.py",
                "blank_guard 命中。跑不通的场景请换成报错终端 / console / "
                "出错源码文件截图，见 SKILL.md 第 5 节。",
            ),
            (
                "evidence_guard.py",
                "evidence_guard 命中。verifiable=是 时必须上传实际页面、"
                "CLI、API 或运行/构建报错证据；如果只能静态走读，请改为不可验证并写原因。",
            ),
            (
                "screenshot_style_guard.py",
                "screenshot_style_guard 命中。验证图不能带黑底/白底边框、header、label "
                "或 padding；单张截图时必须直接使用原图。",
            ),
        ]
        for script, message in guards:
            rc = subprocess.call(
                [sys.executable, str(HERE / script), *args.task_ids],
                cwd=str(REPO_ROOT),
            )
            if rc != 0:
                print(f"==> push 中止：{message}")
                return 1

    summary = []
    for tid in args.task_ids:
        # 已推过的定稿不再重推，除非用户显式删除 sentinel
        pushed_sentinel = REPO_ROOT / "delivery" / tid / ".pushed"
        if not args.dry_run and pushed_to_allowed_sheet(pushed_sentinel):
            print(f"   [skip] {tid}: delivery/{tid}/.pushed 存在（已推过飞书）；"
                  f"如需强推，先手动删除该文件")
            summary.append((tid, 0, 0, "already-pushed"))
            continue
        try:
            t, i = push_one(tid, token, sheet_id, remote_headers,
                            meta["grid_properties"]["row_count"], args.profile, args.dry_run,
                            only_changed_text=args.only_changed_text,
                            only_image_model=args.only_image_model)
            summary.append((tid, t, i, None))
            if not args.dry_run:
                pushed_sentinel.parent.mkdir(parents=True, exist_ok=True)
                pushed_sentinel.write_text(
                    f"pushed via push_feishu.py at row lookup by 任务ID\nsheet_id={sheet_id}\n",
                    encoding="utf-8",
                )
        except Exception as e:
            print(f"   [err] {tid}: {e}")
            summary.append((tid, 0, 0, str(e)))

    print("\n==> 汇总")
    for tid, t, i, err in summary:
        if err:
            print(f"   {tid}: 失败 — {err}")
        else:
            print(f"   {tid}: 文字 {t} 图片 {i}")
    return 0 if all(err is None for _, _, _, err in summary) else 1


if __name__ == "__main__":
    sys.exit(main())
