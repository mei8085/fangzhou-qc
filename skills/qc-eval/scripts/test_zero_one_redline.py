"""Regression tests for the initial-empty-workspace redline delivery path."""

from __future__ import annotations

import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from PIL import Image

import evidence_guard
import generate_model_qc_platform
import push_xlsx
import run
import screenshot_style_guard
import selfcheck
from zero_one_redline import DIRTY_TAG, NOTE, REDLINE, delivery_sheet


MODELS = ("model-a", "model-b", "model-c", "model-d")


class ZeroOneRedlineTests(unittest.TestCase):
    def setUp(self):
        self.qc = {
            "task_id": "example",
            "sheet": {
                "redline_issue": REDLINE,
                "prompt_quality": "高质量",
                "is_high_quality_task": "是",
                "is_normal_task": "是",
                "models": {
                    name: {
                        "rating": "高质量",
                        "note": "旧备注",
                        "execution_process_score": 10,
                        "execution_artifact_score": 10,
                        "audit_reason": "旧审核理由",
                        "dirty_tag": "",
                    }
                    for name in MODELS
                },
            },
        }
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        task_dir = self.root / "tasks" / "example"
        task_dir.mkdir(parents=True)
        (task_dir / "qc_cells.json").write_text(
            json.dumps(self.qc, ensure_ascii=False), encoding="utf-8"
        )

    def workbook(self):
        wb = Workbook()
        ws = wb.active
        headers = [
            "任务质量（题目质量）",
            "题目质量说明",
            "题目红线问题",
            "题目红线备注",
            "是否正常题目",
            "是否是高质量题目",
            "一次质检平台审核理由（按模型）",
        ]
        for name in MODELS:
            headers.extend(
                [
                    f"{name}反馈质量（作业质量）",
                    f"{name}质量备注",
                    f"{name}脏数据标签",
                    f"{name}脏数据标签理由",
                ]
            )
        ws.append(headers)
        with contextlib.redirect_stdout(io.StringIO()):
            push_xlsx.fill_one(ws, selfcheck.find_headers(ws), "example", self.qc, 2)
        return wb

    def test_delivery_overrides_only_the_reviewer_layer(self):
        normalized = delivery_sheet(self.qc["sheet"])
        self.assertEqual(normalized["prompt_quality"], "不合格")
        self.assertEqual(normalized["prompt_quality_note"], NOTE)
        self.assertEqual(normalized["redline_issue"], REDLINE)
        for name in MODELS:
            self.assertEqual(normalized["models"][name]["rating"], "不合格")
            self.assertEqual(normalized["models"][name]["note"], NOTE)
            self.assertEqual(normalized["models"][name]["dirty_tag"], DIRTY_TAG)
        self.assertEqual(self.qc["sheet"]["prompt_quality"], "高质量")
        self.assertEqual(self.qc["sheet"]["models"]["model-a"]["execution_process_score"], 10)

    def test_xlsx_and_platform_payload_use_fixed_zero_one_text(self):
        wb = self.workbook()
        ws = wb.active
        headers = selfcheck.find_headers(ws)
        self.assertEqual(ws.cell(2, headers["任务质量（题目质量）"]).value, "不合格")
        self.assertEqual(ws.cell(2, headers["题目质量说明"]).value, NOTE)
        self.assertEqual(ws.cell(2, headers["model-a质量备注"]).value, NOTE)
        self.assertEqual(ws.cell(2, headers["model-a脏数据标签"]).value, DIRTY_TAG)

        output = generate_model_qc_platform.render("example", self.qc, {})
        self.assertIn("请选择该任务的质量：不可用", output)
        self.assertEqual(output.count("请选择该用户的反馈质量：完全不可用"), 4)
        self.assertEqual(output.count(NOTE), 13)
        self.assertNotIn("旧审核理由", output)

        path = self.root / "result.xlsx"
        wb.save(path)
        with patch.object(selfcheck, "REPO_ROOT", self.root):
            issues, _ = selfcheck.check_xlsx(path)
        self.assertEqual(issues, {})

    def test_empty_materials_do_not_infer_zero_one(self):
        sheet = {"models": {}, "prompt_quality": "不合格"}
        self.assertIs(delivery_sheet(sheet), sheet)

    def test_initial_workspace_evidence_does_not_need_model_map(self):
        shots = self.root / "delivery" / "example" / "screenshots"
        shots.mkdir(parents=True)
        with patch.object(evidence_guard, "REPO_ROOT", self.root), patch.object(
            screenshot_style_guard, "REPO_ROOT", self.root
        ):
            self.assertTrue(evidence_guard.check_task("example"))
            for name in MODELS:
                original = shots / f"{name}_01_initial_workspace_empty.png"
                Image.new("RGB", (320, 240), "white").save(original)
                (shots / f"{name}_verify.png").write_bytes(original.read_bytes())
            self.assertEqual(evidence_guard.check_task("example"), [])
            self.assertEqual(screenshot_style_guard.check_task("example"), [])
            (shots / "model-a_verify.png").write_bytes(b"changed")
            self.assertTrue(evidence_guard.check_task("example"))

    def test_emit_short_circuits_product_detection(self):
        with patch.object(run, "TASKS_DIR", self.root / "tasks"), patch.object(
            run, "pushed_to_current_b_sheet", return_value=False
        ), patch.object(run, "ai_accuracy_enabled", return_value=False), patch.object(
            run, "task_kind", side_effect=AssertionError("must not inspect product")
        ), patch.object(run, "run_py", return_value=0) as execute, contextlib.redirect_stdout(
            io.StringIO()
        ):
            self.assertTrue(
                run.emit_one(
                    "example",
                    False,
                    period="38期",
                    sheet="test",
                    source=Path("source"),
                )
            )
            self.assertEqual(
                [call.args[0] for call in execute.call_args_list],
                [
                    "blank_guard.py",
                    "evidence_guard.py",
                    "screenshot_style_guard.py",
                    "push_xlsx.py",
                    "generate_model_qc_platform.py",
                    "selfcheck.py",
                ],
            )


if __name__ == "__main__":
    unittest.main()
