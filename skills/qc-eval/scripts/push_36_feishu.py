#!/usr/bin/env python3
"""Upload a 38期 delivery workbook to Feishu by guarded 任务ID lookup.

38期 uses the source-sheet `任务ID` as the local workspace/delivery directory
name. The script re-reads the remote sheet, locates the row by `任务ID`, and
refuses to write unless the remote task title matches the local delivery
workbook's task title.

Default target:
  https://bytedance.larkoffice.com/wiki/OY1Dw0dcailCxTk3UGOcykVznAf?sheet=uBlwx4

Usage:
  python skills/qc-eval/scripts/push_38_feishu.py <任务ID> --dry-run
  python skills/qc-eval/scripts/push_38_feishu.py <任务ID> --write --confirm-user-instructed
"""

from __future__ import annotations

import argparse
import subprocess
import sys
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SKILL_DIR = HERE.parent
REPO_ROOT = SKILL_DIR.parent.parent
sys.path.insert(0, str(HERE))

from push_feishu import (  # noqa: E402
    READ_ONLY_HEADERS,
    build_push_targets,
    cell_to_text,
    get_sheet_meta,
    read_range,
    resolve_spreadsheet_token,
    write_cell,
    write_image,
)

DEFAULT_WIKI = "https://bytedance.larkoffice.com/wiki/OY1Dw0dcailCxTk3UGOcykVznAf?sheet=uBlwx4"
DEFAULT_SHEET_ID = "uBlwx4"
DEFAULT_PROFILE = "quan"
ALLOWED_REMOTE_TITLE_PREFIX = "38期"


def read_remote_headers(token: str, sheet_id: str, max_col: int, profile: str) -> dict[str, int]:
    last = get_column_letter(max_col)
    rows = read_range(token, sheet_id, f"A1:{last}3", profile)
    for row in rows:
        headers = {}
        for idx, value in enumerate(row, start=1):
            text = cell_to_text(value).strip()
            if text:
                headers[text] = idx
        if "任务" not in headers and headers.get("任务ID", 0) > 1:
            # Some 38期 exports expose A1 as blank even though it is the task-title
            # column. Keep the guarded task-name check by treating the column before
            # 任务ID as 任务.
            headers["任务"] = headers["任务ID"] - 1
        if "任务" in headers and "任务ID" in headers:
            return headers
    raise RuntimeError("远端前 3 行没有找到「任务」和「任务ID」表头")


def safe_tag(name: str) -> str:
    return re.sub(r"[^0-9A-Za-z_.-]+", "_", name.strip())


def find_remote_row_by_task_id(token: str, sheet_id: str, row_count: int, headers: dict[str, int], task_id: str, profile: str) -> tuple[int | None, str]:
    first = min(headers["任务ID"], headers["任务"])
    last = max(headers["任务ID"], headers["任务"])
    values = read_range(token, sheet_id, f"{get_column_letter(first)}1:{get_column_letter(last)}{row_count}", profile)
    id_idx = headers["任务ID"] - first
    task_idx = headers["任务"] - first
    for offset, row in enumerate(values, start=1):
        if id_idx < len(row) and cell_to_text(row[id_idx]).strip() == task_id:
            title = cell_to_text(row[task_idx]).strip() if task_idx < len(row) else ""
            return offset, title
    return None, ""


def load_delivery(task_id: str, local_xlsx: Path | None) -> tuple[Path, openpyxl.worksheet.worksheet.Worksheet, dict[str, int]]:
    xlsx = local_xlsx or (REPO_ROOT / "delivery" / task_id / f"{task_id}.xlsx")
    if not xlsx.exists():
        raise FileNotFoundError(f"本地交付 xlsx 不存在: {xlsx}")
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.worksheets[0]
    headers = {
        str(ws.cell(1, c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if isinstance(ws.cell(1, c).value, str) and str(ws.cell(1, c).value).strip()
    }
    if "任务" not in headers:
        raise RuntimeError(f"{xlsx} 表头缺少「任务」列")
    return xlsx, ws, headers


def verify_screenshots(task_id: str, skip: bool) -> int:
    if skip:
        return 0
    guards = [
        "blank_guard.py",
        "evidence_guard.py",
        "screenshot_style_guard.py",
    ]
    for script in guards:
        rc = subprocess.call([sys.executable, str(HERE / script), task_id], cwd=str(REPO_ROOT))
        if rc != 0:
            print(f"==> push 中止：{script} 未通过")
            return rc
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("task_id", help="38期源表「任务ID」，也是 tasks/ 和 delivery/ 子目录名")
    parser.add_argument("--wiki-url", default=DEFAULT_WIKI)
    parser.add_argument("--sheet-id", default=DEFAULT_SHEET_ID)
    parser.add_argument("--profile", default=DEFAULT_PROFILE)
    parser.add_argument("--local-xlsx", default=None)
    parser.add_argument("--dry-run", action="store_true", help="只打印计划写入，不实际写飞书")
    parser.add_argument("--write", action="store_true", help="实际写飞书；未指定时默认 dry-run")
    parser.add_argument(
        "--confirm-user-instructed",
        action="store_true",
        help="确认用户已明确要求上传/回写飞书；没有该确认时 --write 会被拒绝",
    )
    parser.add_argument("--only-changed-text", action="store_true", help="只写文字，跳过截图")
    args = parser.parse_args()

    if args.sheet_id != DEFAULT_SHEET_ID:
        print(f"错误：38期只允许上传 sheet={DEFAULT_SHEET_ID}，拒绝 {args.sheet_id}", file=sys.stderr)
        return 2
    if "sheet=" in args.wiki_url and f"sheet={DEFAULT_SHEET_ID}" not in args.wiki_url:
        print(f"错误：38期只允许 URL sheet={DEFAULT_SHEET_ID}", file=sys.stderr)
        return 2
    if args.write and args.dry_run:
        print("错误：--write 和 --dry-run 不能同时使用", file=sys.stderr)
        return 2
    if args.write and not args.confirm_user_instructed:
        print(
            "错误：禁止主动上传飞书。只有用户明确要求上传/回写/写入飞书后，才可使用 "
            "--write --confirm-user-instructed。",
            file=sys.stderr,
        )
        return 2

    dry = not args.write
    local_xlsx = Path(args.local_xlsx) if args.local_xlsx else None
    xlsx, ws, local_headers = load_delivery(args.task_id, local_xlsx)
    if "任务ID" in local_headers:
        local_task_id = str(ws.cell(2, local_headers["任务ID"]).value or "").strip()
        if local_task_id and local_task_id != args.task_id:
            print(f"错误：本地交付任务ID={local_task_id} 与参数 {args.task_id} 不一致", file=sys.stderr)
            return 2
    local_title = str(ws.cell(2, local_headers["任务"]).value or "").strip()
    if not local_title:
        print("错误：本地交付第 2 行「任务」为空", file=sys.stderr)
        return 2

    print(f"==> 38期 Feishu push: task_id={args.task_id} dry-run={dry}")
    print(f"    local xlsx: {xlsx}")
    print(f"    local task: {local_title}")

    token = resolve_spreadsheet_token(args.wiki_url, args.profile)
    meta = get_sheet_meta(token, args.sheet_id, args.profile)
    title = str(meta.get("title", "")).strip()
    if not title.startswith(ALLOWED_REMOTE_TITLE_PREFIX):
        print(f"错误：远端表名不是 38期表，实际为 {title!r}", file=sys.stderr)
        return 2
    row_count = int(meta["grid_properties"]["row_count"])
    remote_headers = read_remote_headers(
        token,
        args.sheet_id,
        int(meta["grid_properties"]["column_count"]),
        args.profile,
    )
    row, remote_title = find_remote_row_by_task_id(token, args.sheet_id, row_count, remote_headers, args.task_id, args.profile)
    if not row:
        print(f"错误：远端表未找到任务ID={args.task_id}", file=sys.stderr)
        return 2
    print(f"    remote task row {row}: {remote_title}")
    if remote_title != local_title:
        print("错误：远端任务ID对应的任务名与本地不一致，拒绝上传", file=sys.stderr)
        return 2

    if verify_screenshots(args.task_id, args.only_changed_text) != 0:
        return 1

    text_headers, image_map = build_push_targets(local_headers)
    text_count = 0
    for header in text_headers:
        if header in READ_ONLY_HEADERS:
            continue
        lcol = local_headers.get(header)
        rcol = remote_headers.get(header)
        if not lcol or not rcol:
            continue
        value = ws.cell(2, lcol).value
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        write_cell(token, args.sheet_id, f"{get_column_letter(rcol)}{row}", str(value), args.profile, dry)
        text_count += 1

    image_count = 0
    if not args.only_changed_text:
        img_dir = REPO_ROOT / "delivery" / args.task_id / "screenshots"
        for header, model in image_map.items():
            rcol = remote_headers.get(header)
            if not rcol:
                continue
            img = img_dir / f"{model.replace(' ', '_')}_verify.png"
            if not img.exists():
                img = img_dir / f"{safe_tag(model)}_verify.png"
            if not img.exists():
                print(f"   [warn] 本地缺图: {img.relative_to(REPO_ROOT)}")
                continue
            write_image(token, args.sheet_id, f"{get_column_letter(rcol)}{row}", img, args.profile, dry)
            image_count += 1

    print(f"==> 完成：文字 {text_count} 单元格，图片 {image_count} 单元格 ({'dry-run' if dry else '已写入'})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
