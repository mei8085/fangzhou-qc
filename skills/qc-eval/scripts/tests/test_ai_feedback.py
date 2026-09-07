import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from check_ai_feedback import check_task, evidence_issues, workbook_labels
from selfcheck import redline_consistency_issues
from openpyxl import Workbook


TAG = "用户问题:AI打分/AI评价"
QUOTE = "本次评分和评价全部由 AI 生成，直接提交。"


class AIFeedbackTests(unittest.TestCase):
    def model(self):
        return {"dirty_tag": TAG, "ai_feedback_evidence": {
            "source": "annotator_feedback", "quote": QUOTE,
            "reason": "标注人员明确承认将 AI 生成的评分和评价直接提交。",
        }}

    def test_task_self_evaluation_is_not_auto_classified(self):
        for text in ["请自我评估并运行测试", "按10分制rubric验收", "自评12/12，PASS", "作为AI，已完成自测"]:
            self.assertEqual([], evidence_issues({"dirty_tag": ""}, {"prompt": text, "output": text}))

    def test_7979_original_wrong_label_is_blocked(self):
        model = {"dirty_tag": TAG, "dirty_tag_reason": "任务要求模型自评，报告包含自评验收结论"}
        self.assertTrue(evidence_issues(model, {"evaluation": {"remark": "测试通过，评分合理"}}))

    def test_task_or_model_sources_cannot_stand_in_for_annotator(self):
        for source in ["task_prompt", "model_output", "execution_report"]:
            model = self.model()
            model["ai_feedback_evidence"]["source"] = source
            self.assertTrue(evidence_issues(model, {"evaluation": {"remark": QUOTE}}))

    def test_quote_must_belong_to_annotator(self):
        self.assertTrue(evidence_issues(self.model(), {"prompt": QUOTE, "output": QUOTE,
                                                      "evaluation": {"remark": "后端测试已复核"}}))

    def test_confirmed_annotator_evidence_is_accepted(self):
        for annotation in [{"remark": QUOTE}, {"summary": QUOTE},
                           {"dimensionResults": [{"remark": QUOTE}]}]:
            self.assertEqual([], evidence_issues(self.model(), {"evaluation": annotation}))

    def test_other_labels_and_format_alone_are_not_ai_evidence(self):
        for tag in ["", "用户问题:敷衍恶意打分", "其他"]:
            self.assertEqual([], evidence_issues({"dirty_tag": tag}, {"evaluation": {"remark": "# 评分 90% ✅"}}))
        model = self.model()
        model["ai_feedback_evidence"]["reason"] = ""
        self.assertTrue(evidence_issues(model, {"evaluation": {"remark": QUOTE}}))

    def test_stale_outgoing_tag_cannot_bypass_evidence_check(self):
        with tempfile.TemporaryDirectory() as tmp:
            task = Path(tmp)
            (task / "qc_cells.json").write_text(json.dumps({"sheet": {"models": {"0825E": {"dirty_tag": ""}}}}))
            (task / "model_map.json").write_text('{"models": {}}')
            self.assertEqual([], check_task(task))
            self.assertTrue(check_task(task, {"0825E": ["用户问题：AI 打分 / AI 评价"]}))

    def test_ai_tag_does_not_require_unrelated_legacy_redline(self):
        ws = Workbook().active
        ws.append(["红线问题", "任务质量（题目质量）", "0825E反馈质量（作业质量）",
                   "0825E质量备注", "0825E脏数据标签"])
        ws.append(["", "合格", "不合格", "标注人员明确承认使用AI生成提交评价", TAG])
        headers = {cell.value: cell.column for cell in ws[1]}
        self.assertEqual(TAG, workbook_labels(ws)["0825E"])
        self.assertEqual([], redline_consistency_issues(ws, headers))
        ws.cell(2, 5).value = "用户问题:敷衍恶意打分"
        self.assertIn("redline-rating-missing-红线问题", redline_consistency_issues(ws, headers))


if __name__ == "__main__":
    unittest.main()
