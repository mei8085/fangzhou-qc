#!/usr/bin/env python3
"""Probe 38期 Feishu task IDs and compare with a local xlsx snapshot.

38期必须使用源表「任务ID」作为本地工作目录名；远端行号只作为诊断信息，
不能作为 tasks/ 或 delivery/ 子目录名。

This script is read-only for Feishu. It reuses the repo's lark-cli wrapper from
push_feishu.py; it does not use feishu-docs.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from lib.paths import default_source_xlsx  # noqa: E402
from push_feishu import (  # noqa: E402
    cell_to_text,
    get_sheet_meta,
    read_range,
    resolve_spreadsheet_token,
)

DEFAULT_WIKI = "https://bytedance.larkoffice.com/wiki/OY1Dw0dcailCxTk3UGOcykVznAf?sheet=uBlwx4"
DEFAULT_SHEET_ID = "uBlwx4"
DEFAULT_PROFILE = "quan"
DEFAULT_LOCAL = default_source_xlsx()


def find_header(values: list[list]) -> tuple[int, dict[str, int]]:
    for r_idx, row in enumerate(values[:5], start=1):
        headers: dict[str, int] = {}
        for c_idx, value in enumerate(row, start=1):
            text = cell_to_text(value).strip()
            if text:
                headers[text] = c_idx
        if "任务" in headers:
            return r_idx, headers
    raise RuntimeError("远端前 5 行没有找到「任务」表头")


def row_value(row: list, col: int) -> str:
    return cell_to_text(row[col - 1]).strip() if col - 1 < len(row) else ""


def load_local_tasks(local_path: Path, sheet_name: str | None) -> tuple[str, int, dict[int, str]]:
    if not local_path.exists():
        return "", 0, {}
    wb = openpyxl.load_workbook(local_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    task_col = None
    header_row = 0
    for r in range(1, min(ws.max_row, 5) + 1):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(r, c).value or "").strip() == "任务":
                task_col = c
                header_row = r
                break
        if task_col:
            break
    if not task_col:
        return ws.title, 0, {}
    tasks: dict[int, str] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        tasks[r] = str(ws.cell(r, task_col).value or "").strip()
    return ws.title, header_row, tasks


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--wiki-url", default=DEFAULT_WIKI)
    parser.add_argument("--sheet-id", default=DEFAULT_SHEET_ID)
    parser.add_argument("--profile", default=DEFAULT_PROFILE)
    parser.add_argument("--local", default=str(DEFAULT_LOCAL), help="local xlsx snapshot to compare")
    parser.add_argument("--local-sheet", default=None)
    parser.add_argument("--task-title", default=None)
    parser.add_argument("--out", default="tasks/_38_task_ids.json")
    args = parser.parse_args()

    token = resolve_spreadsheet_token(args.wiki_url, args.profile)
    meta = get_sheet_meta(token, args.sheet_id, args.profile)
    row_count = meta["grid_properties"]["row_count"]
    col_count = meta["grid_properties"]["column_count"]
    last = get_column_letter(col_count)
    values = read_range(token, args.sheet_id, f"A1:{last}{row_count}", args.profile)
    header_row, headers = find_header(values)
    task_col = headers["任务"]
    task_id_col = headers.get("任务ID")

    rows = []
    for row_no in range(header_row + 1, len(values) + 1):
        row = values[row_no - 1]
        task = row_value(row, task_col)
        if not task:
            continue
        rows.append({
            "row": row_no,
            "task": task,
            "task_id": row_value(row, task_id_col) if task_id_col else "",
        })

    local_title, local_header_row, local_tasks = load_local_tasks(Path(args.local), args.local_sheet)
    mismatches = []
    matched = 0
    compared = 0
    for item in rows:
        remote_task = item["task"]
        local_task = local_tasks.get(item["row"], "")
        if not local_task and not remote_task:
            continue
        compared += 1
        if local_task == remote_task:
            matched += 1
        else:
            mismatches.append({
                "row": item["row"],
                "remote_task": remote_task,
                "local_task": local_task,
            })

    target_matches = []
    if args.task_title:
        target_matches = [r for r in rows if r["task"] == args.task_title]

    by_title: dict[str, list[int]] = {}
    for item in rows:
        by_title.setdefault(item["task"], []).append(item["row"])
    duplicate_titles = {title: row_nos for title, row_nos in by_title.items() if len(row_nos) > 1}
    blank_task_ids = sum(1 for item in rows if not item["task_id"])

    out = {
        "wiki_url": args.wiki_url,
        "spreadsheet_token": token,
        "sheet_id": args.sheet_id,
        "sheet_title": meta.get("title"),
        "row_count": row_count,
        "column_count": col_count,
        "header_row": header_row,
        "headers": headers,
        "task_rows": rows,
        "workspace_id_rule": "use source sheet 任务ID only",
        "blank_task_id_count": blank_task_ids,
        "duplicate_titles": duplicate_titles,
        "local": {
            "path": str(Path(args.local)),
            "sheet": local_title,
            "header_row": local_header_row,
            "compared": compared,
            "matched": matched,
            "mismatch_count": len(mismatches),
            "mismatches_sample": mismatches[:30],
        },
        "target_matches": target_matches,
    }
    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"sheet={meta.get('title')} id={args.sheet_id} rows={row_count} cols={col_count}")
    print(f"header_row={header_row} task_col={task_col} task_id_col={task_id_col or '-'}")
    print(f"remote task rows={len(rows)}")
    print(f"blank task_id rows={blank_task_ids}/{len(rows)}")
    print(f"duplicate titles={len(duplicate_titles)}")
    if args.task_title:
        print(f"target matches={len(target_matches)}")
        for item in target_matches:
            print(f"  row {item['row']} task_id={item['task_id'] or '-'}: {item['task']}")
    if local_tasks:
        print(f"local={args.local} sheet={local_title} header_row={local_header_row}")
        print(f"same-row compare: matched={matched}/{compared}, mismatches={len(mismatches)}")
        for item in mismatches[:10]:
            print(f"  row {item['row']}: remote={item['remote_task']!r} local={item['local_task']!r}")
    else:
        print(f"local={args.local}: no comparable task column found")
    print(f"wrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
