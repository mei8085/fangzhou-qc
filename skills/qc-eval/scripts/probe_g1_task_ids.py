#!/usr/bin/env python3
"""Probe G1-new tasks from the local source xlsx (docs/G1-new.xlsx).

G1-new does not sync the QC working set from Feishu. The local workbook is
the source of truth; Feishu sheet=cHr9N5 is only the delivery target.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from lib.g1 import (  # noqa: E402
    CODE_TO_PLATFORM,
    PERIOD,
    SHEET_MODELS,
    SOURCE_SHEET,
    SOURCE_XLSX,
)
from lib.paths import default_g1_source_xlsx  # noqa: E402


def cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_rows(path: Path, sheet_name: str | None) -> tuple[str, int, dict[str, int], list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row = 0
    headers: dict[str, int] = {}
    for r in range(1, min(ws.max_row, 5) + 1):
        found: dict[str, int] = {}
        for c in range(1, ws.max_column + 1):
            text = cell_text(ws.cell(r, c).value)
            if text:
                found[text] = c
        if "任务ID" in found:
            header_row = r
            headers = found
            break
    if not header_row:
        raise RuntimeError(f"{path} 前 5 行没有「任务ID」表头")

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        task_id = cell_text(ws.cell(r, headers["任务ID"]).value)
        task = cell_text(ws.cell(r, headers.get("任务", 0)).value) if headers.get("任务") else ""
        if not task_id and not task:
            continue
        rows.append({
            "row": r,
            "task_id": task_id,
            "task": task,
            "user": cell_text(ws.cell(r, headers["做题用户"]).value) if "做题用户" in headers else "",
            "reviewer": cell_text(ws.cell(r, headers["质检人"]).value) if "质检人" in headers else "",
            "supplier": cell_text(ws.cell(r, headers["所属商"]).value) if "所属商" in headers else "",
            "label": cell_text(ws.cell(r, headers["label"]).value) if "label" in headers else "",
        })
    title = ws.title
    wb.close()
    return title, header_row, headers, rows


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", default=str(default_g1_source_xlsx()))
    parser.add_argument("--sheet", default=SOURCE_SHEET)
    parser.add_argument("--task-id", default=None, help="只打印该任务ID")
    parser.add_argument("--task-title", default=None, help="只打印该任务名")
    parser.add_argument("--out", default="tasks/_g1_task_ids.json")
    args = parser.parse_args()

    source = Path(args.source)
    if not source.exists():
        print(f"ERROR: G1-new 源表不存在: {source}", file=sys.stderr)
        return 2

    sheet_title, header_row, headers, rows = load_rows(source, args.sheet)
    matches = rows
    if args.task_id:
        matches = [r for r in rows if r["task_id"] == str(args.task_id).strip()]
    if args.task_title:
        matches = [r for r in matches if r["task"] == args.task_title]

    payload = {
        "period": PERIOD,
        "source": str(source),
        "sheet": sheet_title,
        "header_row": header_row,
        "headers": headers,
        "sheet_models": SHEET_MODELS,
        "platform_map": CODE_TO_PLATFORM,
        "task_rows": rows,
        "target_matches": matches,
        "workspace_id_rule": "use local source 任务ID; same ID is used to search Ark project 277",
    }
    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"period={PERIOD} source={source} sheet={sheet_title} header_row={header_row}")
    print(f"task rows={len(rows)}")
    print("models:")
    for code, platform in CODE_TO_PLATFORM.items():
        print(f"  {code}/{platform} -> 【G1】{code}")
    if args.task_id or args.task_title:
        print(f"target matches={len(matches)}")
        if not matches:
            print("ERROR: 本地源表未找到该任务", file=sys.stderr)
            print(f"wrote {out_path}")
            return 1
        for item in matches:
            print(
                f"  row {item['row']} task_id={item['task_id']}: {item['task']} "
                f"user={item['user'] or '-'} reviewer={item['reviewer'] or '-'}"
            )
    else:
        for item in rows:
            print(f"  {item['task_id']}\t{item['task']}")
    print(f"wrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
