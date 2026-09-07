#!/usr/bin/env python3
"""Self-check reviewer-owned cells in delivery/<TASK_ID>/<TASK_ID>.xlsx.

38期反馈质量评的是标注人的打分、备注、佐证与 trace 是否可信，不是直接评
模型作业本身。命中任一项即视为不合格，调用方负责改写后重跑 push_xlsx.py。

Exit code: 0=clean, 1=has issues.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent.parent.parent
sys.path.insert(0, str(HERE))
from lib.annotator_tags import (  # noqa: E402
    iter_model_evaluations,
    qc_coverage_issues,
    self_test as annotator_tag_self_test,
)
from lib.writing_style import (  # noqa: E402
    consecutive_opening_issues,
    dangling_opening_issues,
    load_writing_style,
    opening_family_example_issues,
    repeated_audit_ending_issues,
    scan_banned_style,
    scan_feedback_rating_rationale,
)
from check_prompt_cheat import build_notes  # noqa: E402
from lib.g1 import (  # noqa: E402
    FEISHU_DIRTY_TAG_OPTIONS,
    MISSING_MODEL_RATING,
    PERIOD,
    PLATFORM_CODES,
    SHEET_MODELS,
    is_empty_feishu_dirty_tag,
    is_feishu_dirty_tag_header,
    sheet_model_label,
    to_feishu_dirty_tag,
)
from lib.label_options import LABEL_HEADER, label_options  # noqa: E402
from lib.prompt_cheat import is_prompt_cheat_task  # noqa: E402

# ---------- 表头 → reviewer 列定位 ----------

_TEXT_HEADER_KEYS = [
    "题目质量说明",
    "不可验证原因",
    "0614C对比GLM 5.2 模型差异（轨迹/产物）pattern简述",
    "红线问题",
    "题目红线问题",
    "题目红线备注",
    "质检人",
    "所属商",
    "所述商",
    "机器质检结果是否正确",
    "错误原因",
]
_MODEL_NOTE_SUFFIX = "质量备注"
_MODEL_RATING_SUFFIX = "反馈质量（作业质量）"
_HIGH_QUALITY_TASK_HEADER = "是否是高质量题目"
_NORMAL_TASK_HEADER = "是否正常题目"
_CATEGORY_CORRECT_HEADER = "分类是否正确"
_TASK_RATING_HEADER = "任务质量（题目质量）"
_REDLINE_HEADER = "红线问题"
_G1_REDLINE_HEADER = "题目红线问题"
_MACHINE_QC_HEADER = "机器质检结果是否正确"
_MACHINE_QC_REASON_HEADER = "错误原因"
_PLATFORM_AUDIT_HEADER = "一次质检平台审核理由（按模型）"
_PROMPT_QUALITY_NOTE_HEADER = "题目质量说明"
_HIGH_QUALITY_FEEDBACK_VALUE = "高质量"
_CURRENT_PERIODS = {PERIOD}

TASK_REDLINE_OPTIONS = {
    "毫无意义的题目",
    "非代码题",
    "题目不合理",
    "非本人场景",
    "静态展示性页面",
}
FEEDBACK_REDLINE_OPTIONS = {
    "任务质量（题目质量）触及红线",
    "虚假反馈",
    "trace和反馈结果明显不一致",
    "完全敷衍",
}
REDLINE_OPTIONS = TASK_REDLINE_OPTIONS | FEEDBACK_REDLINE_OPTIONS
REDLINE_ALIAS_MAP = {
    "任务质量触及红线": "任务质量（题目质量）触及红线",
    "trace 和反馈结果明显不一致": "trace和反馈结果明显不一致",
}


def find_headers(ws) -> dict[str, int]:
    """delivery xlsx 的表头在第 1 行（rebuild 后固定）。"""
    return {
        str(ws.cell(1, c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if isinstance(ws.cell(1, c).value, str) and ws.cell(1, c).value.strip()
    }


def detect_period(headers: dict[str, int]) -> str:
    if any(h.startswith("【G1】") for h in headers) or "题目红线问题" in headers:
        return "G1-new"
    if (
        "一次质检平台审核理由（按模型）" in headers
        and any(h.endswith(_MODEL_RATING_SUFFIX) for h in headers)
    ):
        return PERIOD
    return "38期"


def reviewer_columns(headers: dict[str, int]) -> tuple[list[int], set[int]]:
    """Return (all_reviewer_cols, model_note_cols).

    Model-note columns are those ending with `质量备注` (每个模型一列).
    """
    cols: list[int] = []
    model_note_cols: set[int] = set()
    for h, c in headers.items():
        if h in _TEXT_HEADER_KEYS:
            cols.append(c)
        elif h.endswith(_MODEL_NOTE_SUFFIX):
            cols.append(c)
            model_note_cols.add(c)
    return sorted(cols), model_note_cols


# ---------- 规则 ----------

# 38期反馈质量评的是标注人对模型产出结果的打分与备注。
# 模型备注必须出现下面任一关键词，说明这段在评"标注人评价是否准确"。
ANNOTATOR_REF = re.compile(
    r"标注人|打分|备注|分值|\d+\s*分|识别(?:到|出|准确)|判断(?:准确|对|错|有偏差)|佐证"
)
FORBIDDEN_WORDS = [
    "自动化", "脚本", "pipeline", "compose", "screenshot_vue",
    "按 spec", "spec 维度", "spec 第", "稽核", "实测属实",
    "SKILL.md", "qc_cells", "model_map",
    "——", "—",
    "【", "】", "「", "」", "『", "』",
    # 把模型称作"这家"/"四家"是口语错觉（像在评供应商）。规范统一写作
    # "这个模型"/"该模型"/"多个模型"。
    "这家", "那家", "几家", "两家", "三家", "四家", "五家",
    # "封顶"是方舟平台打分时的内部勾选项名，对题面下游 reviewer 不通用，
    # 备注里要翻译成"认为核心功能不可用 / 严重问题"这种白话。
    "封顶",
    # 当前反馈质量只判断打分是否准确、备注是否清楚、证据是否支撑结论；
    # 不再评价"打分格式规范"，避免把旧模板句误写进交付备注。
    "打分及打分备注均未按照打分格式规范进行",
    "未按照打分格式规范",
    "打分格式规范",
    "按评分项规范填写",
]

# 交付描述里不要出现 AI 腔短评和档位套话。完整词表见
# skills/qc-eval/config/writing_style.json；这里保留一份兜底，避免配置缺失时漏拦。
UNNATURAL_STYLE_WORDS = [
    "更稳", "方向到位", "观察到位", "覆盖到位", "佐证到位",
    "定性正确", "抓得准", "判断落点", "描述属实", "贴合",
    "属实", "靠谱", "偏薄", "兜底",
    "站得住", "高一档", "低一档", "判断成立", "标签成立",
    "评分方向合理", "完全吻合", "方向基本准确",
    "说得通", "对得上", "对不上", "能对上", "全对上",
    "写重了", "抬上去", "抬在", "没有抬", "过重",
    "同向", "偏松", "偏严", "缺口", "偏重", "未写反", "说得过头", "仍在",
    "可核", "也可核", "能核",
    "，与备注相符", "。与备注相符", "，与反馈相符", "。与反馈相符",
]
UNNATURAL_STYLE_RE = re.compile(
    r"(?:很|挺|蛮|够|更|比较|较|不太|不够)准(?!确)|"
    r"(?:判断|识别|方向)准(?!确)|抓得(?:很)?准(?!确)|"
    r"(?:判断|识别)?落点"
)
SCORE_IN_NOTE = re.compile(r"\d+\s*分")
MISSING_MODEL_NOTE_RE = re.compile(r"未作答|未测试|未提交|未完成.{0,8}(?:模型|测试)|没有运行记录")
MODEL_IDENTIFIER_TERMS = sorted(
    {str(v).strip() for v in [*SHEET_MODELS, *PLATFORM_CODES] if str(v).strip()},
    key=len,
    reverse=True,
)
MODEL_IDENTIFIER_RE = re.compile(
    r"(?<![A-Za-z0-9_-])(?:"
    + "|".join(re.escape(v) for v in MODEL_IDENTIFIER_TERMS)
    + r")(?![A-Za-z0-9_-])",
    re.IGNORECASE,
) if MODEL_IDENTIFIER_TERMS else None

FORBIDDEN_SYM = re.compile(r"[*#]|`")

# 句首或紧跟标点 / 空白后的"我"。
FIRST_PERSON = re.compile(r"(?:^|[，。、；！？\s])我(?![们国方司家])")

# 单元格内代码引用过多
CODE_REF = re.compile(r"\.[a-z]{1,5}(?=\b|[　-鿿])|src/|第\s*\d+\s*行")

# 备注末尾不要重复定档
TRAILING_VERDICT = re.compile(r"(?:给|定|定档为?|评)\s*[高中低](?![分级数等数])")

MAX_LEN = 220
MAX_CODE_REFS = 3

# 题目质量说明是给表格 reviewer 看的判断依据，只说明题目本身的完整性、
# 经济价值、评测价值、可复核性和未达高质量原因；不要把“是否是高质量题目”
# 这个表字段的填写动作写进去。涉及 Coding Agent workspace 规模时只写
# 概括判断，不暴露具体代码量阈值或代码行数。区分度从题目设计简洁说明即可，
# 不要求展开决策空间等细节，也不能引用测试模型实际分数、维度分或分差。
# 页面上看不到的聚合分由
# writing_style.json 的 page_hidden_* 统一拦截。
TASK_NOTE_FORBIDDEN_PATTERNS = [
    (re.compile(r"是否是高质量题目.{0,20}填"), "task-note-no-high-quality-field-self-description"),
    (re.compile(r"填[“\"]?[是否][”\"]?"), "task-note-no-field-fill-instruction"),
    (re.compile(r"(?:\b\d+(?:\.\d+)?\s*k\+?(?=\W|$)|\d{1,3}(?:,\d{3})+\s*\+?\s*行|\d{3,6}\s*\+?\s*行)", re.IGNORECASE), "task-note-no-specific-code-volume-threshold"),
    (re.compile(r"(?:代码|仓库|workspace).{0,12}(?:约|大约|大概|不足|不到|少于|低于|只有|仅有)?\s*\d{2,6}\s*(?:行|lines?)", re.IGNORECASE), "task-note-no-specific-code-line-count"),
    (re.compile(r"\d{2,6}\s*(?:行|lines?).{0,12}(?:代码|仓库|workspace)", re.IGNORECASE), "task-note-no-specific-code-line-count"),
    (re.compile(r"(?:执行过程|过程|产物质量|产物)\s*分(?!析)"), "task-note-no-test-model-dimension-scores"),
    (re.compile(r"模型.{0,24}(?:实际得分|评分结果|分差|相差\s*\d+\s*分)"), "task-note-no-test-model-score-difference"),
    (
        re.compile(
            r"(?:" + "|".join(re.escape(v) for v in MODEL_IDENTIFIER_TERMS) + r").{0,12}\d+\s*分",
            re.IGNORECASE,
        ),
        "task-note-no-test-model-scores",
    ),
]

# 复测不能把质检机本地缺依赖当作模型产物缺陷。备注里如果出现典型本地
# 环境失败信号，必须同时说明已经补齐环境并重跑后的结果。
LOCAL_ENV_FAILURE_RE = re.compile(
    r"(?:本地|测试机?|质检机?).{0,8}(?:环境|依赖).{0,12}(?:缺少|缺失|未安装|没装|失败|不可用)|"
    r"(?:缺少|缺失|未安装|没装).{0,12}(?:依赖|环境|Playwright|Chromium|Electron|node_modules)|"
    r"(?:command not found|Cannot find module|ModuleNotFoundError|No module named|"
    r"No matching version found|Electron failed to install|Playwright .*browser|"
    r"(?:vite|tsc|vitest|playwright): command not found)",
    re.IGNORECASE,
)
ENV_RECHECK_RE = re.compile(
    r"(?:环境|依赖|Playwright|Chromium|Electron).{0,16}(?:补齐|补装|安装|重装|修复|恢复)|"
    r"(?:补齐|补装|安装|重装|修复|恢复).{0,16}(?:环境|依赖|Playwright|Chromium|Electron)|"
    r"(?:npm|pnpm|yarn|pip|playwright|electron).{0,24}(?:install|安装)|"
    r"(?:重跑|复跑|重新(?:验证|测试|运行)|安装后|补齐后)",
    re.IGNORECASE,
)

# 标注人备注若带有明确环境线索，而质检复测环境不同或未确认相同环境，
# 不能直接否定标注人的失败归因；必须写成“本次环境未复现/无法排除环境差异”。
ANNOTATOR_ENV_MARKER_RE = re.compile(
    r"Windows|macOS|Mac\s?OS|Linux|Ubuntu|CI|Docker|Chromium|Chrome|"
    r"Playwright|Node(?:\.js)?\s*\d|Python\s*\d|C:\\|C:/|Program Files|"
    r"PATH|shell|端口|localhost|127\.0\.0\.1|本地环境|测试机|浏览器版本|"
    r"空格路径|大小写敏感|换行符",
    re.IGNORECASE,
)
ASSERTIVE_ENV_MISMATCH_RE = re.compile(
    r"(?:标注人?(?:记录|备注|写|认为)?.{0,60}(?:E2E|Playwright|测试|构建|启动|运行|报错|失败|Failed to fetch).{0,40}(?:不符|不一致|错误|错))|"
    r"(?:(?:失败|原因|理由|归因|记录).{0,24}(?:不符|不一致|错误|错))|"
    r"(?:与(?:实测|实际情况|验证结果)不符)",
    re.IGNORECASE,
)
ENV_DIFF_CAUTION_RE = re.compile(
    r"本次环境|当前环境|本轮环境|本次复测环境|未复现|结果不同|无法排除环境差异|"
    r"不能排除环境差异|可能存在环境差异|环境差异|未在.{0,12}(?:Windows|macOS|Linux|CI|Docker)|"
    r"(?:同环境|相同环境).{0,16}(?:复测|验证|复现)|结论(?:仍)?可采信|仅部分可采信",
    re.IGNORECASE,
)

# 红线必须有明确证据；出现这类模糊措辞时按"疑罪从无"拦下。
VAGUE_REDLINE_EVIDENCE = re.compile(
    r"疑似|可能|不确定|不好说|无法确认|猜测|推测|大概|也许|似乎|像是|感觉"
)
MIN_REDLINE_EVIDENCE_LEN = 20


def scan_style_only(s: str) -> list[str]:
    issues: list[str] = []
    issues.extend(scan_banned_style(s))
    for w in UNNATURAL_STYLE_WORDS:
        token = f"style:{w}"
        if w in s and token not in issues:
            issues.append(token)
    for m in UNNATURAL_STYLE_RE.findall(s):
        token = f"style:{m}"
        if token not in issues:
            issues.append(token)
    return issues


def scan_cell(
    s: str,
    *,
    col: int,
    model_note_cols: set[int],
    period: str,
    skip_annotator_checks: bool = False,
) -> list[str]:
    issues = []
    for w in FORBIDDEN_WORDS:
        if w in s:
            issues.append(f"word:{w}")
    issues.extend(scan_style_only(s))
    if FORBIDDEN_SYM.search(s):
        issues.append("forbidden-sym")
    if FIRST_PERSON.search(s):
        issues.append("first-person:我")
    if len(s) > MAX_LEN:
        issues.append(f"too-long:{len(s)}")
    if len(CODE_REF.findall(s)) > MAX_CODE_REFS:
        issues.append("too-many-code-refs")
    if TRAILING_VERDICT.search(s):
        issues.append("trailing-verdict")
    if col in model_note_cols and not skip_annotator_checks:
        if period in _CURRENT_PERIODS:
            if "标注" not in s and "轮次作弊" not in s:
                issues.append("not-evaluating-annotator")
        elif not ANNOTATOR_REF.search(s):
            issues.append("not-evaluating-annotator")
        if period in _CURRENT_PERIODS and not SCORE_IN_NOTE.search(s) and not MISSING_MODEL_NOTE_RE.search(s):
            issues.append("model-note-missing-score")
        if period in _CURRENT_PERIODS and MODEL_IDENTIFIER_RE and MODEL_IDENTIFIER_RE.search(s):
            issues.append("model-note-no-model-identifier")
        issues.extend(dangling_opening_issues(s))
    return issues


def scan_task_quality_note(s: str) -> list[str]:
    issues: list[str] = []
    for pattern, issue in TASK_NOTE_FORBIDDEN_PATTERNS:
        if pattern.search(s):
            issues.append(issue)
    return issues


def check_platform_qc_md(md_path: Path) -> list[str]:
    if not md_path.exists():
        return []
    issues: list[str] = []
    label = md_path.name
    for idx, line in enumerate(md_path.read_text(encoding="utf-8-sig").splitlines(), start=1):
        found = scan_style_only(line)
        if found:
            issues.append(f"{label}:{idx}: {found}")
    return issues


def _notes_from_xlsx(xlsx: Path) -> tuple[str, list[str]]:
    """Return (题目质量说明, 各模型质量备注) from a delivery workbook."""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active
    headers = find_headers(ws)
    _, model_note_cols = reviewer_columns(headers)
    task_col = headers.get(_PROMPT_QUALITY_NOTE_HEADER)
    task_note = cell_text(ws, task_col)
    model_notes = [cell_text(ws, col) for col in sorted(model_note_cols)]
    return task_note, model_notes


def _notes_from_qc_cells(path: Path) -> tuple[str, list[str]]:
    """Return (题目质量说明, 各模型质量备注) from qc_cells.json."""
    data = json.loads(path.read_text(encoding="utf-8"))
    sheet = data.get("sheet") or data
    task_note = str(sheet.get("prompt_quality_note") or "")
    models = sheet.get("models") or {}
    model_notes: list[str] = []
    if isinstance(models, dict):
        for item in models.values():
            if isinstance(item, dict):
                model_notes.append(str(item.get("note") or ""))
    return task_note, model_notes


def load_recent_task_notes(
    task_id: str, window: int
) -> list[tuple[str, str, list[str]]]:
    """Return newest-first ``(task_id, task_note, model_notes)`` for recent tasks."""
    found: dict[str, tuple[float, str, list[str]]] = {}
    delivery = REPO_ROOT / "delivery"
    if delivery.is_dir():
        for folder in delivery.iterdir():
            if not folder.is_dir() or folder.name == task_id:
                continue
            xlsx = folder / f"{folder.name}.xlsx"
            if not xlsx.exists():
                continue
            try:
                task_note, model_notes = _notes_from_xlsx(xlsx)
            except Exception:
                continue
            found[folder.name] = (xlsx.stat().st_mtime, task_note, model_notes)
    tasks = REPO_ROOT / "tasks"
    if tasks.is_dir():
        for folder in tasks.iterdir():
            if not folder.is_dir() or folder.name == task_id or folder.name in found:
                continue
            qc_path = folder / "qc_cells.json"
            if not qc_path.exists():
                continue
            try:
                task_note, model_notes = _notes_from_qc_cells(qc_path)
            except Exception:
                continue
            found[folder.name] = (qc_path.stat().st_mtime, task_note, model_notes)
    ranked = sorted(found.items(), key=lambda item: item[1][0], reverse=True)
    return [(tid, payload[1], payload[2]) for tid, payload in ranked[:window]]


def split_redline_options(s: str) -> list[str]:
    return [p.strip() for p in re.split(r"[，,；;\n]+", s) if p.strip()]


def canonical_redline_option(s: str) -> str | None:
    if s in REDLINE_OPTIONS:
        return s
    if s in REDLINE_ALIAS_MAP:
        return REDLINE_ALIAS_MAP[s]
    compact = re.sub(r"\s+", "", s)
    if compact in REDLINE_OPTIONS:
        return compact
    if compact in REDLINE_ALIAS_MAP:
        return REDLINE_ALIAS_MAP[compact]
    return None


def cell_text(ws, col: int | None) -> str:
    if not col:
        return ""
    return str(ws.cell(2, col).value or "").strip()


def is_redline_rating(ws, col: int | None) -> bool:
    return "红线" in cell_text(ws, col)


def can_carry_feedback_redline(ws, col: int | None) -> bool:
    """38期模型反馈质量下拉只有 高质量/合格/不合格。

    反馈红线标签放在「红线问题」列；对应模型反馈质量应为「不合格」。
    兼容历史本地文件里曾出现的「红线」值。
    """
    value = cell_text(ws, col)
    return "红线" in value or value == "不合格"


def can_carry_task_redline(ws, col: int | None) -> bool:
    """Task-quality 不合格 can carry 题目红线 tags such as 题目不合理."""
    value = cell_text(ws, col)
    return "红线" in value or value == "不合格"


def model_rating_note_cols(headers: dict[str, int]) -> list[tuple[str, int, int | None]]:
    out = []
    for h, rating_col in headers.items():
        if not h.endswith(_MODEL_RATING_SUFFIX):
            continue
        model = h[:-len(_MODEL_RATING_SUFFIX)]
        out.append((model, rating_col, headers.get(f"{model}{_MODEL_NOTE_SUFFIX}")))
    return out


def missing_model_qc_issues(ws, headers: dict[str, int], task_id: str) -> dict[int, list[str]]:
    """Require absent Ark models to be explicitly marked unqualified."""
    if not task_id:
        return {}
    map_path = REPO_ROOT / "tasks" / task_id / "model_map.json"
    if not map_path.exists():
        return {}
    try:
        model_map = json.loads(map_path.read_text(encoding="utf-8-sig"))
    except Exception:
        return {}
    columns = {model: (rating_col, note_col) for model, rating_col, note_col in model_rating_note_cols(headers)}
    bad: dict[int, list[str]] = {}
    for model, info in (model_map.get("models") or {}).items():
        if not info.get("missing"):
            continue
        rating_col, note_col = columns.get(model, (None, None))
        if not rating_col:
            continue
        rating = cell_text(ws, rating_col)
        note = cell_text(ws, note_col)
        dirty_col = headers.get(f"{model}脏数据标签")
        dirty = cell_text(ws, dirty_col)
        if rating != MISSING_MODEL_RATING:
            bad.setdefault(rating_col, []).append(f"missing-model-rating-must-be-{MISSING_MODEL_RATING}")
        if not MISSING_MODEL_NOTE_RE.search(note):
            bad.setdefault(note_col or rating_col, []).append("missing-model-note-must-state-not-submitted")
        if dirty != "用户问题:敷衍恶意打分":
            bad.setdefault(dirty_col or rating_col, []).append("missing-model-dirty-tag-must-be-用户问题:敷衍恶意打分")
    return bad


def evidence_too_thin(texts: list[str]) -> bool:
    concrete = [t for t in texts if t.strip()]
    if not concrete:
        return True
    return sum(len(t.strip()) for t in concrete) < MIN_REDLINE_EVIDENCE_LEN


def redline_header_col(headers: dict[str, int]) -> int | None:
    return headers.get(_G1_REDLINE_HEADER) or headers.get(_REDLINE_HEADER)


def is_artifact_download_unavailable_case(qc: dict | None) -> bool:
    sheet = ((qc or {}).get("sheet") or {})
    return sheet.get("artifact_download_unavailable") is True


def redline_consistency_issues(
    ws,
    headers: dict[str, int],
    qc: dict | None = None,
) -> list[str]:
    """Check 红线问题 options, category match, and minimum evidence discipline."""
    redline_col = redline_header_col(headers)
    if not redline_col:
        return []

    selected = str(ws.cell(2, redline_col).value or "").strip()
    raw_options = split_redline_options(selected)
    canonical_options = [canonical_redline_option(v) for v in raw_options]
    valid_options = [v for v in canonical_options if v]

    task_rating_redline = is_redline_rating(ws, headers.get(_TASK_RATING_HEADER))
    task_can_carry_redline = can_carry_task_redline(ws, headers.get(_TASK_RATING_HEADER))
    task_note = cell_text(ws, headers.get("题目质量说明"))
    redline_model_notes = [
        cell_text(ws, note_col)
        for _model, rating_col, note_col in model_rating_note_cols(headers)
        if can_carry_feedback_redline(ws, rating_col)
    ]
    # Historical task 6699 establishes a non-redline exception: when every
    # model artifact download is disabled, feedback is unqualified because it
    # cannot be verified, while the dirty-data tag is "其他" and redline stays
    # empty. A dedicated guard below validates that exception in full.
    has_model_redline = bool(redline_model_notes) and not is_artifact_download_unavailable_case(qc)
    has_redline_rating = task_rating_redline or has_model_redline

    issues: list[str] = []
    if has_redline_rating and not selected:
        issues.append("redline-rating-missing-红线问题")
    if selected:
        invalid = [raw for raw, canon in zip(raw_options, canonical_options) if not canon]
        if invalid:
            issues.append(f"invalid-红线问题:{'，'.join(invalid)}")
        if not has_redline_rating:
            issues.append("红线问题-without-redline-rating")
        if VAGUE_REDLINE_EVIDENCE.search(selected):
            issues.append("redline-evidence-vague")

    selected_set = set(valid_options)
    has_task_tag = bool(selected_set & TASK_REDLINE_OPTIONS)
    has_feedback_tag = bool(selected_set & FEEDBACK_REDLINE_OPTIONS)
    if task_rating_redline and selected and not has_task_tag:
        issues.append("task-redline-missing-task-tag")
    if has_model_redline and selected and not has_feedback_tag and not has_task_tag:
        issues.append("feedback-redline-missing-feedback-tag")

    for opt in selected_set:
        if opt in TASK_REDLINE_OPTIONS and not task_can_carry_redline:
            issues.append(f"task-redline-tag-without-task-redline-rating:{opt}")
        elif opt == "任务质量（题目质量）触及红线" and not task_rating_redline:
            issues.append("feedback-task-quality-tag-without-task-redline")
        elif opt in FEEDBACK_REDLINE_OPTIONS and opt != "任务质量（题目质量）触及红线" and not has_model_redline:
            issues.append(f"feedback-redline-tag-without-model-redline-rating:{opt}")

    if has_task_tag:
        if evidence_too_thin([task_note]):
            issues.append("task-redline-evidence-too-thin")
        elif VAGUE_REDLINE_EVIDENCE.search(task_note):
            issues.append("task-redline-evidence-vague")
    feedback_tags_needing_model_evidence = selected_set & (
        FEEDBACK_REDLINE_OPTIONS - {"任务质量（题目质量）触及红线"}
    )
    if feedback_tags_needing_model_evidence:
        if evidence_too_thin(redline_model_notes):
            issues.append("feedback-redline-evidence-too-thin")
        elif VAGUE_REDLINE_EVIDENCE.search("\n".join(redline_model_notes)):
            issues.append("feedback-redline-evidence-vague")
    if "任务质量（题目质量）触及红线" in selected_set:
        if evidence_too_thin([task_note]):
            issues.append("task-quality-trigger-redline-evidence-too-thin")
        elif VAGUE_REDLINE_EVIDENCE.search(task_note):
            issues.append("task-quality-trigger-redline-evidence-vague")
    return issues


def artifact_download_unavailable_issues(
    ws,
    headers: dict[str, int],
    qc: dict | None,
    task_id: str,
) -> dict[int, list[str]]:
    """Validate the explicit all-download-buttons-disabled exception."""
    if not is_artifact_download_unavailable_case(qc):
        return {}
    bad: dict[int, list[str]] = {}

    def add(col: int | None, issue: str) -> None:
        if col:
            bad.setdefault(col, []).append(issue)

    verifiable_col = headers.get("题目是否可验证") or headers.get("题目是否可验")
    reason_col = headers.get("不可验证原因")
    redline_col = redline_header_col(headers)
    if cell_text(ws, verifiable_col) != "否":
        add(verifiable_col, "download-unavailable-verifiable-must-be-否")
    reason = cell_text(ws, reason_col)
    if "产物不可下载" not in reason or "无法验证" not in reason:
        add(reason_col, "download-unavailable-reason-must-name-artifact-and-verification")
    if cell_text(ws, redline_col):
        add(redline_col, "download-unavailable-redline-must-be-empty")

    qc_models = (((qc or {}).get("sheet") or {}).get("models") or {})
    for model, rating_col, note_col in model_rating_note_cols(headers):
        dirty_col = headers.get(f"{model}脏数据标签")
        dirty_reason_col = headers.get(f"{model}脏数据标签理由")
        if cell_text(ws, rating_col) != "不合格":
            add(rating_col, "download-unavailable-rating-must-be-不合格")
        note = cell_text(ws, note_col)
        if "产物不可下载" not in note or "无法验证" not in note:
            add(note_col or rating_col, "download-unavailable-note-must-name-artifact-and-verification")
        if cell_text(ws, dirty_col) != "其他":
            add(dirty_col or rating_col, "download-unavailable-dirty-tag-must-be-其他")
        dirty_reason = cell_text(ws, dirty_reason_col)
        if "产物不可下载" not in dirty_reason or "无法验证" not in dirty_reason:
            add(dirty_reason_col or rating_col, "download-unavailable-dirty-reason-invalid")
        audit = str((qc_models.get(model) or {}).get("platform_audit_reason") or "")
        if "产物不可下载" not in audit or "无法验证" not in audit or "完全不可用" not in audit:
            add(rating_col, "download-unavailable-platform-audit-reason-invalid")

        screenshots = REPO_ROOT / "delivery" / task_id / "screenshots"
        if not (screenshots / f"{model}_01_download_blocker.png").exists():
            add(note_col or rating_col, "download-unavailable-evidence-screenshot-missing")
        if not (screenshots / f"{model}_verify.png").exists():
            add(note_col or rating_col, "download-unavailable-verify-screenshot-missing")

    discovery_path = REPO_ROOT / "tasks" / task_id / "ark_g1" / "download_list_discover.json"
    try:
        discovery = json.loads(discovery_path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        add(reason_col, "download-unavailable-discovery-missing-or-invalid")
    else:
        if discovery.get("classification") != "artifact_download_unavailable":
            add(reason_col, "download-unavailable-discovery-classification-invalid")
        if discovery.get("all_expected_models_present") is not True:
            add(reason_col, "download-unavailable-discovery-models-incomplete")
        if discovery.get("all_download_buttons_disabled") is not True:
            add(reason_col, "download-unavailable-discovery-buttons-not-all-disabled")
    return bad


def machine_qc_issues(ws, headers: dict[str, int]) -> dict[int, list[str]]:
    correct_col = headers.get(_MACHINE_QC_HEADER)
    if not correct_col:
        return {}
    reason_col = headers.get(_MACHINE_QC_REASON_HEADER)
    verdict = cell_text(ws, correct_col)
    reason = cell_text(ws, reason_col)
    bad: dict[int, list[str]] = {}
    if verdict not in {"正确", "部分正确可参考", "错误"}:
        bad.setdefault(correct_col, []).append("machine-qc-verdict-must-be-正确-部分正确可参考-or-错误")
    if verdict in {"部分正确可参考", "错误"} and len(reason) < 15:
        bad.setdefault(reason_col or correct_col, []).append("machine-qc-mismatch-reason-too-thin")
    if verdict == "正确" and not reason:
        bad.setdefault(reason_col or correct_col, []).append("machine-qc-correct-reason-missing")
    return bad


def high_quality_task_issues(ws, headers: dict[str, int]) -> dict[int, list[str]]:
    col = headers.get(_HIGH_QUALITY_TASK_HEADER)
    if not col:
        return {}
    value = str(ws.cell(2, col).value or "").strip()
    if not value:
        return {col: ["required-empty"]}
    if value not in {"是", "否"}:
        return {col: [f"invalid-option:{value}"]}
    model_ratings = [
        cell_text(ws, rating_col)
        for _model, rating_col, _note_col in model_rating_note_cols(headers)
        if cell_text(ws, rating_col)
    ]
    task_rating = cell_text(ws, headers.get(_TASK_RATING_HEADER))
    if task_rating and model_ratings:
        all_high = task_rating == _HIGH_QUALITY_FEEDBACK_VALUE and all(
            r == _HIGH_QUALITY_FEEDBACK_VALUE for r in model_ratings
        )
        expected = "是" if all_high else "否"
        if value != expected:
            return {
                col: [
                    f"must-match-task-and-feedback-quality:{expected}",
                    f"task-rating:{task_rating}",
                    f"model-ratings:{'/'.join(model_ratings)}",
                ]
            }
    return {}


def normal_task_issues(ws, headers: dict[str, int]) -> dict[int, list[str]]:
    col = headers.get(_NORMAL_TASK_HEADER)
    if not col:
        return {}
    value = str(ws.cell(2, col).value or "").strip()
    if not value:
        return {col: ["required-empty"]}
    if value not in {"是", "否"}:
        return {col: [f"invalid-option:{value}"]}
    return {}


def category_correct_issues(ws, headers: dict[str, int]) -> dict[int, list[str]]:
    col = headers.get(_CATEGORY_CORRECT_HEADER)
    if not col:
        return {}
    value = str(ws.cell(2, col).value or "").strip()
    if not value:
        return {col: ["required-empty"]}
    if value not in {"是", "否"}:
        return {col: [f"invalid-option:{value}"]}
    return {}


def label_issues(ws, headers: dict[str, int]) -> dict[int, list[str]]:
    col = headers.get(LABEL_HEADER)
    if not col:
        return {}
    value = str(ws.cell(2, col).value or "").strip()
    if not value:
        return {col: ["required-empty"]}
    allowed = set(label_options())
    if allowed and value not in allowed:
        return {col: [f"invalid-option:{value}"]}
    return {}


def g1_dirty_tag_issues(ws, headers: dict[str, int], period: str) -> dict[int, list[str]]:
    """G1 Feishu dirty-tag cells must be empty or an exact dropdown option."""
    if period not in _CURRENT_PERIODS:
        return {}
    allowed = set(FEISHU_DIRTY_TAG_OPTIONS)
    bad: dict[int, list[str]] = {}
    for header, col in headers.items():
        if not is_feishu_dirty_tag_header(header):
            continue
        raw = cell_text(ws, col)
        if raw == "非脏数据":
            bad[col] = ["feishu-dirty-tag-leave-empty"]
            continue
        if is_empty_feishu_dirty_tag(raw):
            continue
        mapped = to_feishu_dirty_tag(raw)
        parts = [p for p in mapped.split("、") if p]
        invalid = [p for p in parts if p not in allowed]
        if mapped != raw or invalid:
            bad[col] = [f"invalid-feishu-dirty-tag:{raw}"]
    return bad


def audit_line_for_model(combined: str, model: str) -> str:
    """Return the per-model sentence from the combined platform-audit cell."""
    label = sheet_model_label(model)
    prefixes = (f"{label}：", f"{label}:", f"{model}：", f"{model}:")
    for line in str(combined or "").splitlines():
        text = line.strip()
        for prefix in prefixes:
            if text.startswith(prefix):
                return text[len(prefix):].strip()
    return ""


def feedback_rating_rationale_issues(
    ws,
    headers: dict[str, int],
    qc: dict | None,
    *,
    prompt_cheat: bool,
) -> dict[int, list[str]]:
    """Require Feishu notes to name 高质量/合格/不合格; audits to name 完全可用/部分可用/完全不可用."""
    combined = cell_text(ws, headers.get(_PLATFORM_AUDIT_HEADER))
    qc_models = ((qc or {}).get("sheet") or {}).get("models") or {}
    bad: dict[int, list[str]] = {}
    for model, rating_col, note_col in model_rating_note_cols(headers):
        rating = cell_text(ws, rating_col)
        note = cell_text(ws, note_col)
        data = qc_models.get(model) or {}
        audit = str(
            data.get("platform_audit_reason") or data.get("audit_reason") or ""
        ).strip() or audit_line_for_model(combined, model)
        issues = scan_feedback_rating_rationale(
            note,
            audit,
            rating,
            prompt_cheat=prompt_cheat,
        )
        if not issues:
            continue
        target = note_col or rating_col
        bad.setdefault(target, []).extend(issues)
        if any(item.startswith("audit-") for item in issues) and headers.get(
            _PLATFORM_AUDIT_HEADER
        ):
            bad.setdefault(headers[_PLATFORM_AUDIT_HEADER], []).extend(
                [item for item in issues if item.startswith("audit-")]
            )
    return bad


def local_env_dependency_issues(ws, headers: dict[str, int]) -> dict[int, list[str]]:
    """Block conclusions that treat unremediated local dependency failures as product bugs."""
    bad: dict[int, list[str]] = {}
    cols: set[int] = set()
    _, model_note_cols = reviewer_columns(headers)
    cols.update(c for c in model_note_cols if c)
    for header in ("不可验证原因", _PLATFORM_AUDIT_HEADER):
        col = headers.get(header)
        if col:
            cols.add(col)
    for col in sorted(cols):
        text = cell_text(ws, col)
        if not text:
            continue
        if LOCAL_ENV_FAILURE_RE.search(text) and not ENV_RECHECK_RE.search(text):
            bad.setdefault(col, []).append("local-env-dependency-not-rechecked")
    return bad


def annotator_environment_attribution_issues(
    ws,
    headers: dict[str, int],
    qc: dict | None,
    task_id: str,
) -> dict[int, list[str]]:
    """Block overconfident contradiction claims when annotator env differs.

    If the annotator explicitly names an environment-sensitive context
    (Windows, CI, Docker, Playwright version, local ports, path format, etc.),
    a QC note may say the current run did not reproduce the failure, but it
    must not directly call the annotator's failure reason wrong unless the note
    also documents same-environment proof or environment-difference caution.
    """
    if not task_id:
        return {}
    rows = iter_model_evaluations(task_id)
    if not rows:
        return {}
    qc_models = ((qc or {}).get("sheet") or {}).get("models") or {}
    combined = cell_text(ws, headers.get(_PLATFORM_AUDIT_HEADER))
    note_cols = {model: note_col for model, _, note_col in model_rating_note_cols(headers)}
    bad: dict[int, list[str]] = {}

    for row in rows:
        sheet = str(row.get("sheet") or "")
        annotator_text = " ".join(
            str(part or "")
            for dim in row.get("dimensions") or []
            for part in (dim.get("remark"), dim.get("overall_remark"))
        )
        if not ANNOTATOR_ENV_MARKER_RE.search(annotator_text):
            continue

        note_col = note_cols.get(sheet)
        note = cell_text(ws, note_col) if note_col else ""
        data = qc_models.get(sheet) or {}
        audit = str(
            data.get("platform_audit_reason") or data.get("audit_reason") or ""
        ).strip() or audit_line_for_model(combined, sheet)

        for label, text, target_col in (
            ("note", note, note_col),
            ("audit", audit, headers.get(_PLATFORM_AUDIT_HEADER)),
        ):
            if not text:
                continue
            if not ASSERTIVE_ENV_MISMATCH_RE.search(text):
                continue
            if ENV_DIFF_CAUTION_RE.search(text):
                continue
            if target_col:
                bad.setdefault(target_col, []).append(
                    f"{label}-env-diff-overconfident-annotator-mismatch"
                )
    return bad


def annotator_tag_qc_issues(
    ws,
    headers: dict[str, int],
    qc: dict | None,
    task_id: str,
    *,
    prompt_cheat: bool,
) -> dict[int, list[str]]:
    """Fail when defect tags contradict remarks but QC did not cover that."""
    if prompt_cheat or not task_id:
        return {}
    models = iter_model_evaluations(task_id)
    if not models:
        return {}
    qc_models = ((qc or {}).get("sheet") or {}).get("models") or {}
    combined = cell_text(ws, headers.get(_PLATFORM_AUDIT_HEADER))
    bad: dict[int, list[str]] = {}
    rating_cols = {model: rating_col for model, rating_col, _ in model_rating_note_cols(headers)}
    note_cols = {model: note_col for model, _, note_col in model_rating_note_cols(headers)}
    for row in models:
        sheet = str(row.get("sheet") or "")
        rating_col = rating_cols.get(sheet)
        note_col = note_cols.get(sheet)
        rating = cell_text(ws, rating_col) if rating_col else ""
        note = cell_text(ws, note_col) if note_col else ""
        data = qc_models.get(sheet) or {}
        audit = str(
            data.get("platform_audit_reason") or data.get("audit_reason") or ""
        ).strip() or audit_line_for_model(combined, sheet)
        issues = qc_coverage_issues(row, rating, note, audit)
        if not issues:
            continue
        target = note_col or rating_col
        if target:
            bad.setdefault(target, []).extend(issues)
    return bad


def load_qc_cells(task_id: str) -> dict | None:
    """Return tasks/<TASK_ID>/qc_cells.json when the file exists."""
    if not task_id:
        return None
    path = REPO_ROOT / "tasks" / task_id / "qc_cells.json"
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception:
        return None


def task_upload_collection_issues(task_id: str) -> list[str]:
    """Require current-period first-round uploads before delivery."""
    ark_dir = REPO_ROOT / "tasks" / task_id / "ark_g1"
    task_path = ark_dir / "task.json"
    manifest_path = ark_dir / "task_uploads.json"
    if not task_path.exists():
        return ["task-upload-task-json-missing"]
    try:
        task = json.loads(task_path.read_text(encoding="utf-8-sig"))
    except Exception:
        return ["task-upload-task-json-invalid"]
    attachments = task.get("attachments")
    if attachments is None and isinstance(task.get("task"), dict):
        attachments = task["task"].get("attachments")
    if not isinstance(attachments, list):
        return ["task-upload-attachments-field-missing"]
    if not manifest_path.exists():
        return ["task-upload-manifest-missing"]
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    except Exception:
        return ["task-upload-manifest-invalid"]
    issues: list[str] = []
    if manifest.get("status") != "complete":
        issues.append(f"task-upload-status:{manifest.get('status')}")
    if int(manifest.get("reported_count") or 0) != len(attachments):
        issues.append("task-upload-reported-count-mismatch")
    if len(manifest.get("items") or []) != len(attachments):
        issues.append("task-upload-downloaded-count-mismatch")
    for key in ("inspection_path", "prompt_materials_path"):
        rel_path = str(manifest.get(key) or "")
        if not rel_path or not (ark_dir / rel_path).exists():
            issues.append(f"task-upload-{key.replace('_', '-')}-missing")
    prompt_report = REPO_ROOT / "tasks" / task_id / "prompt_cheat.json"
    if not prompt_report.exists():
        issues.append("task-upload-prompt-scan-missing")
    else:
        try:
            report = json.loads(prompt_report.read_text(encoding="utf-8-sig"))
        except Exception:
            issues.append("task-upload-prompt-scan-invalid")
        else:
            expected = int(manifest.get("unique_text_material_count") or 0)
            got = int(report.get("uploaded_text_source_count") or 0)
            if expected != got:
                issues.append("task-upload-prompt-source-count-mismatch")
    return issues


def check_xlsx(xlsx_path: Path, task_id: str = "") -> tuple[dict[int, list[str]], str]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    headers = find_headers(ws)
    period = detect_period(headers)
    cols, model_note_cols = reviewer_columns(headers)
    task_note_col = headers.get(_PROMPT_QUALITY_NOTE_HEADER)
    qc = load_qc_cells(task_id)
    prompt_cheat = bool(task_id) and is_prompt_cheat_task(REPO_ROOT / "tasks" / task_id, qc)
    bad: dict[int, list[str]] = {}
    for c in cols:
        v = ws.cell(2, c).value
        if not v:
            continue
        text = str(v)
        issues = scan_cell(
            text,
            col=c,
            model_note_cols=model_note_cols,
            period=period,
            skip_annotator_checks=prompt_cheat,
        )
        if c == task_note_col:
            issues.extend(scan_task_quality_note(text))
        if issues:
            bad[c] = issues
    model_notes = [cell_text(ws, c) for c in sorted(model_note_cols)]
    variety_anchor = min(model_note_cols) if model_note_cols else None
    platform_audit_col = headers.get(_PLATFORM_AUDIT_HEADER)
    platform_audit_text = cell_text(ws, platform_audit_col)
    if platform_audit_col and platform_audit_text:
        for issue in scan_style_only(platform_audit_text):
            bad.setdefault(platform_audit_col, []).append(issue)
        if not prompt_cheat:
            for issue in repeated_audit_ending_issues(platform_audit_text):
                bad.setdefault(platform_audit_col, []).append(issue)
        if "任务质量高" in platform_audit_text:
            bad.setdefault(platform_audit_col, []).append("platform-audit-no-任务质量高")
    if task_id and not prompt_cheat:
        cfg = load_writing_style()
        window = int(
            cfg.get("consecutive_task_opening_window")
            or cfg.get("consecutive_task_opening_window")
            or 3
        )
        recent = load_recent_task_notes(task_id, window)
        if variety_anchor and model_notes:
            for issue in consecutive_opening_issues(
                model_notes,
                [(prev_id, notes) for prev_id, _, notes in recent],
                cfg,
                kind="note",
            ):
                bad.setdefault(variety_anchor, []).append(issue)
        if task_note_col:
            task_note = cell_text(ws, task_note_col)
            for issue in consecutive_opening_issues(
                [task_note],
                [(prev_id, [note]) for prev_id, note, _ in recent],
                cfg,
                kind="task",
            ):
                bad.setdefault(task_note_col, []).append(issue)
    redline_col = redline_header_col(headers)
    redline_issues = redline_consistency_issues(ws, headers, qc)
    if redline_col and redline_issues:
        bad.setdefault(redline_col, []).extend(redline_issues)
    for col, issues in high_quality_task_issues(ws, headers).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in normal_task_issues(ws, headers).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in category_correct_issues(ws, headers).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in label_issues(ws, headers).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in g1_dirty_tag_issues(ws, headers, period).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in machine_qc_issues(ws, headers).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in artifact_download_unavailable_issues(
        ws, headers, qc, task_id
    ).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in feedback_rating_rationale_issues(
        ws, headers, qc, prompt_cheat=prompt_cheat
    ).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in missing_model_qc_issues(ws, headers, task_id).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in local_env_dependency_issues(ws, headers).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in annotator_environment_attribution_issues(
        ws, headers, qc, task_id
    ).items():
        bad.setdefault(col, []).extend(issues)
    for col, issues in annotator_tag_qc_issues(
        ws, headers, qc, task_id, prompt_cheat=prompt_cheat
    ).items():
        bad.setdefault(col, []).extend(issues)
    return bad, period


def writing_template_issues() -> list[str]:
    """Canned examples in writing_style.json and cheat notes must pass rationale rules."""
    issues = opening_family_example_issues()
    bad_task_note = "多个模型的执行过程分和产物质量分均仅相差 1 分，未达到区分度要求。"
    if not scan_task_quality_note(bad_task_note):
        issues.append("task-note-model-score-guard-missing")
    good_task_note = "但题目设计与校验标准未达到高质量题目的区分度参考要求，因此题目质量为合格。"
    if scan_task_quality_note(good_task_note):
        issues.append("task-note-design-rationale-false-positive")
    bad_process_note = (
        "复测显示页面流程、章节加载和响应式验证均通过，"
        "标注人给执行过程 7 分偏低，反馈质量为合格。"
    )
    bad_process_audit = "页面流程通过，执行过程 7 分偏低，反馈部分可用。"
    bad_process_issues = scan_feedback_rating_rationale(
        bad_process_note, bad_process_audit, "合格"
    )
    if not any("process-score-gap" in item for item in bad_process_issues):
        issues.append("process-score-scope-guard-missing")
    reviewer_recheck_note = (
        "补齐依赖后测试通过，标注人给执行过程 6 分偏低；"
        "产物质量 5 分相符，反馈质量为合格。"
    )
    reviewer_recheck_audit = "补齐依赖后测试通过，过程 6 分偏低，反馈部分可用。"
    reviewer_recheck_issues = scan_feedback_rating_rationale(
        reviewer_recheck_note, reviewer_recheck_audit, "合格"
    )
    if not any("process-score-uses-unattributed-validation" in item for item in reviewer_recheck_issues):
        issues.append("unattributed-validation-process-scope-guard-missing")
    model_validation_note = (
        "测试模型在轨迹中运行 go test 后继续排查失败，过程 6 分相符；"
        "产物质量 5 分偏高，反馈质量为合格。"
    )
    model_validation_audit = "模型自测后的排查不足，产物 5 分偏高，反馈部分可用。"
    model_validation_issues = scan_feedback_rating_rationale(
        model_validation_note, model_validation_audit, "合格"
    )
    if any("process-score-uses-reviewer-recheck" in item for item in model_validation_issues):
        issues.append("model-owned-validation-process-scope-false-positive")
    for item in annotator_tag_self_test():
        issues.append(f"annotator-tag-detector:{item}")
    _, specs = build_notes("工具轮次达到 450")
    for idx, (note, audit) in enumerate(specs):
        for item in scan_feedback_rating_rationale(
            note, audit, "不合格", prompt_cheat=True
        ):
            issues.append(f"cheat-template-{idx}:{item}")
    return issues


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("task_ids", nargs="+")
    args = p.parse_args()
    total_bad = 0
    template_issues = writing_template_issues()
    if template_issues:
        print("  [BAD] writing templates")
        for issue in template_issues:
            print(f"     {issue}")
        total_bad += 1
    for tid in args.task_ids:
        xlsx = REPO_ROOT / "delivery" / tid / f"{tid}.xlsx"
        if not xlsx.exists():
            print(f"  [skip] {tid}: {xlsx} not found")
            continue
        bad, period = check_xlsx(xlsx, tid)
        md_issues = check_platform_qc_md(REPO_ROOT / "delivery" / tid / "platform_qc.md")
        md_issues.extend(
            check_platform_qc_md(REPO_ROOT / "delivery" / tid / f"{tid}_model_qc_platform.md")
        )
        if period == PERIOD:
            md_issues.extend(task_upload_collection_issues(tid))
        if bad or md_issues:
            print(f"  [BAD] {tid}  ({period})")
            for col, issues in bad.items():
                print(f"     col {col}: {issues}")
            for issue in md_issues:
                print(f"     platform_qc.md {issue}")
            total_bad += 1
        else:
            print(f"  [OK]  {tid}  ({period})")
    return 0 if total_bad == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
