"""Read and merge the supplemental AI trajectory-QC workbook.

This module owns only the supplemental-data boundary. The regular accuracy
calculation and platform/Feishu update flows remain in their existing scripts.
"""

from __future__ import annotations

import copy
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from .machine_qc import has_returned_machine_qc_result


REQUIRED_HEADERS = (
    "project_id",
    "task_id",
    "run_id",
    "task_title",
    "qc_task_status",
    "qc_task_quality",
    "qc_task_review_quality",
    "qc_task_settlement",
    "qc_task_is_dirty",
    "qc_task_dirty_labels",
    "qc_task_hitl",
    "qc_traj_status",
    "qc_traj_grade",
    "qc_traj_is_dirty",
    "qc_traj_dirty_labels",
    "qc_traj_res_detail",
    "qc_task_res_detail",
)

VALID_GRADES = {"高质量", "合格", "不合格"}


def cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def json_value(value: Any, *, default: Any) -> Any:
    if value is None or value == "":
        return copy.deepcopy(default)
    if not isinstance(value, str):
        return copy.deepcopy(value)
    try:
        return json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError(f"补充表 JSON 字段解析失败：{value[:120]!r}") from exc


def bool_value(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return cell_text(value).lower() in {"true", "1", "yes", "y", "是"}


def first_row_headers(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for index, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=False)), start=1):
        header = cell_text(cell.value)
        if header:
            headers[header] = index
    missing = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing:
        raise ValueError("补充表缺少字段：" + "、".join(missing))
    return headers


@dataclass(frozen=True)
class SupplementRecord:
    task_id: str
    run_id: str
    project_id: str
    task_title: str
    task_status: str
    task_quality: str
    task_review_quality: str
    task_settlement: str
    task_is_dirty: bool
    task_dirty_labels: list[Any]
    task_hitl: list[Any]
    trajectory_status: str
    trajectory_grade: str
    trajectory_is_dirty: bool
    trajectory_dirty_labels: list[Any]
    trajectory_reason: str
    trace_stats: dict[str, Any]
    algorithm_version: str

    @classmethod
    def from_row(cls, row: tuple[Any, ...], headers: dict[str, int]) -> "SupplementRecord":
        def get(name: str) -> Any:
            return row[headers[name] - 1] if headers[name] - 1 < len(row) else None

        trajectory_detail = json_value(get("qc_traj_res_detail"), default={})
        if not isinstance(trajectory_detail, dict):
            raise ValueError("qc_traj_res_detail 必须是 JSON 对象")
        feedback_quality = trajectory_detail.get("feedback_quality") or {}
        if not isinstance(feedback_quality, dict):
            feedback_quality = {}
        grade = cell_text(get("qc_traj_grade")) or cell_text(feedback_quality.get("grade"))
        if grade not in VALID_GRADES:
            raise ValueError(f"run_id={cell_text(get('run_id'))} 的 AI 质检结论无效：{grade!r}")
        status = cell_text(get("qc_traj_status"))
        if status != "success":
            raise ValueError(f"run_id={cell_text(get('run_id'))} 的 AI 质检状态不是 success：{status!r}")

        try:
            task_detail = json_value(get("qc_task_res_detail"), default={})
        except ValueError:
            # Historical task-level records can be truncated by the source
            # export. The trajectory detail is the authoritative field for
            # this supplement, so ignore only the optional task-level blob.
            task_detail = {}
        if not isinstance(task_detail, dict):
            task_detail = {}
        algorithm_version = (
            cell_text(trajectory_detail.get("algorithm_version"))
            or cell_text(task_detail.get("algorithm_version"))
            or "coding-trajectory-v4-2026-08-19"
        )
        dirty_labels = json_value(get("qc_traj_dirty_labels"), default=[])
        if not dirty_labels:
            dirty_labels = trajectory_detail.get("dirty_labels") or []
        if not isinstance(dirty_labels, list):
            dirty_labels = [dirty_labels]
        task_dirty_labels = json_value(get("qc_task_dirty_labels"), default=[])
        if not isinstance(task_dirty_labels, list):
            task_dirty_labels = [task_dirty_labels]
        hitl = json_value(get("qc_task_hitl"), default=[])
        if not isinstance(hitl, list):
            hitl = [hitl]
        trace_stats = trajectory_detail.get("trace_stats") or {}
        if not isinstance(trace_stats, dict):
            trace_stats = {}

        return cls(
            task_id=cell_text(get("task_id")),
            run_id=cell_text(get("run_id")),
            project_id=cell_text(get("project_id")),
            task_title=cell_text(get("task_title")),
            task_status=cell_text(get("qc_task_status")),
            task_quality=cell_text(get("qc_task_quality")),
            task_review_quality=cell_text(get("qc_task_review_quality")),
            task_settlement=cell_text(get("qc_task_settlement")),
            task_is_dirty=bool_value(get("qc_task_is_dirty")),
            task_dirty_labels=task_dirty_labels,
            task_hitl=hitl,
            trajectory_status=status,
            trajectory_grade=grade,
            trajectory_is_dirty=bool_value(get("qc_traj_is_dirty")),
            trajectory_dirty_labels=dirty_labels,
            trajectory_reason=cell_text(feedback_quality.get("reason")),
            trace_stats=trace_stats,
            algorithm_version=algorithm_version,
        )

    def result_json(self) -> dict[str, Any]:
        return {
            "Issues": [],
            "Reason": self.trajectory_reason,
            "IsDirty": self.trajectory_is_dirty,
            "TraceStats": self.trace_stats,
            "DirtyLabels": self.trajectory_dirty_labels,
            "FeedbackGrade": self.trajectory_grade,
        }


def load_supplement_records(source: Path, task_id: str) -> list[SupplementRecord]:
    if not source.exists():
        raise FileNotFoundError(f"AIQC 补充表不存在：{source}")
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if not workbook.worksheets:
        raise ValueError(f"AIQC 补充表没有工作表：{source}")
    worksheet = workbook.worksheets[0]
    headers = first_row_headers(worksheet)
    records: list[SupplementRecord] = []
    seen_runs: set[str] = set()
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        if cell_text(row[headers["task_id"] - 1]) != str(task_id):
            continue
        record = SupplementRecord.from_row(row, headers)
        if record.run_id in seen_runs:
            raise ValueError(f"AIQC 补充表存在重复 run_id：{record.run_id}")
        seen_runs.add(record.run_id)
        records.append(record)
    if not records:
        raise ValueError(f"AIQC 补充表未找到 task_id={task_id}")
    task_ids = {record.task_id for record in records}
    if task_ids != {str(task_id)}:
        raise ValueError(f"AIQC 补充表任务 ID 不一致：{sorted(task_ids)}")
    return records


def trajectory_is_complete(audit: dict[str, Any]) -> bool:
    return has_returned_machine_qc_result(audit)


def merge_audit_trajectory(
    audit: dict[str, Any],
    record: SupplementRecord,
    *,
    force: bool = False,
) -> tuple[dict[str, Any], str]:
    if trajectory_is_complete(audit) and not force:
        return audit, "skipped_complete"

    updated = copy.deepcopy(audit)
    ai_qc = updated.setdefault("aiQc", {})
    current = ai_qc.get("trajectoryResult")
    if not isinstance(current, dict):
        current = {}
    trajectory = copy.deepcopy(current)
    trajectory["status"] = record.trajectory_status
    trajectory["algorithmVersion"] = record.algorithm_version
    trajectory["resultJson"] = record.result_json()
    ai_qc["trajectoryResult"] = trajectory
    return updated, "applied"


def new_audit_cache(
    task_id: str,
    model_row: dict[str, Any],
    record: SupplementRecord,
) -> dict[str, Any]:
    """Build a minimal audit cache when collection did not save audit.json."""
    run_id = cell_text(model_row.get("task_model_run_id") or model_row.get("run_id"))
    try:
        task_model_run_id: Any = int(run_id)
    except ValueError:
        task_model_run_id = run_id
    try:
        numeric_task_id: Any = int(task_id)
    except ValueError:
        numeric_task_id = task_id
    try:
        numeric_project_id: Any = int(record.project_id)
    except ValueError:
        numeric_project_id = record.project_id
    return {
        "taskModelRunId": task_model_run_id,
        "taskId": numeric_task_id,
        "projectId": numeric_project_id,
        "aiQc": {
            "trajectoryResult": {
                "id": None,
                "status": record.trajectory_status,
                "algorithmVersion": record.algorithm_version,
                "resultJson": record.result_json(),
            },
            "feedback": None,
        },
    }


def merge_task_summary(
    task_payload: dict[str, Any],
    records: list[SupplementRecord],
    *,
    force: bool = False,
) -> tuple[dict[str, Any], str]:
    if not records:
        return task_payload, "skipped_no_records"
    current = task_payload.get("aiQc")
    if isinstance(current, dict) and current.get("status") == "success" and not force:
        return task_payload, "skipped_complete"

    first = records[0]
    task_fields = [
        (
            record.task_status,
            record.task_quality,
            record.task_review_quality,
            record.task_settlement,
            record.task_is_dirty,
            record.task_dirty_labels,
            record.task_hitl,
        )
        for record in records
    ]
    if any(fields != task_fields[0] for fields in task_fields[1:]):
        raise ValueError("同一任务的 AIQC 补充记录包含不一致的任务级结果")
    updated = copy.deepcopy(task_payload)
    ai_qc = copy.deepcopy(current) if isinstance(current, dict) else {}
    ai_qc.update(
        {
            "status": first.task_status,
            "taskQuality": first.task_quality,
            "expertEvaluationQuality": first.task_review_quality,
            "isDirty": first.task_is_dirty,
            "settlementRecommendation": first.task_settlement,
            "reviewRoute": "manual_review" if first.task_hitl else "automatic",
            "algorithmVersion": "coding-task-v3-2026-08-19",
        }
    )
    updated["aiQc"] = ai_qc
    return updated, "applied"
